"""Собирает docs/documentation.docx из docs/documentation.md.

Запуск из корня репозитория: python docs/build_documentation.py
Нужен python-docx. Оглавление Word заполняет при открытии или при выгрузке в PDF
(docs/build_documentation_pdf.ps1 обновляет поля и сохраняет PDF рядом).
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

DOCS = Path(__file__).resolve().parent
SOURCE = DOCS / "documentation.md"
TARGET = DOCS / "documentation.docx"

FONT = "Calibri"
MONO = "Consolas"
ACCENT = RGBColor(0x0B, 0x4A, 0xA2)
MUTED = RGBColor(0x5B, 0x64, 0x70)
TEXT_WIDTH_CM = 17.0


def shade(element, fill: str) -> None:
    props = element.get_or_add_tcPr() if element.tag.endswith("}tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def field(paragraph, instruction: str, placeholder: str = "") -> None:
    """Поле Word: номер страницы или оглавление."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run._r.append(begin)
    run = paragraph.add_run()
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    run._r.append(text)
    run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    paragraph.add_run(placeholder)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def inline(paragraph, text: str, size: float | None = None, bold: bool = False) -> None:
    """Текст с `кодом`: код моноширинным шрифтом."""
    for i, part in enumerate(re.split(r"`([^`]*)`", text)):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if i % 2:
            run.font.name = MONO
            run.font.size = Pt((size or 10.5) - 1)


def setup(doc: Document) -> None:
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    section.top_margin = section.bottom_margin = Cm(2.0)
    section.left_margin = section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before in (("Heading 1", 18, 0), ("Heading 2", 13, 12)):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT
        fonts = style.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(footer, "PAGE", "1")
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("Расклад. Сопроводительная документация")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    section.different_first_page_header_footer = True


def title_page(doc: Document, title: str, lines: list[str]) -> None:
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(title.split(". ")[0])
    run.font.size = Pt(40)
    run.bold = True
    run.font.color.rgb = ACCENT
    p = doc.add_paragraph()
    run = p.add_run(title.split(". ", 1)[1])
    run.font.size = Pt(22)
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        inline(p, line, size=11.5)
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph()
    run = p.add_run("Содержание")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = ACCENT
    field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u', "Оглавление обновится при открытии в Word")


def code_block(doc: Document, lines: list[str]) -> None:
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(4 if i == 0 else 0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Cm(0.3)
        shade(p._p, "F1F4F8")
        run = p.add_run(line if line else " ")
        run.font.name = MONO
        run.font.size = Pt(8.5)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(8)


def column_widths(rows: list[list[str]]) -> list[float]:
    count = len(rows[0])
    longest = [max(min(len(r[c]) if c < len(r) else 0, 60) for r in rows) + 6 for c in range(count)]
    total = sum(longest)
    return [TEXT_WIDTH_CM * n / total for n in longest]


def table(doc: Document, rows: list[list[str]]) -> None:
    widths = column_widths(rows)
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for r, cells in enumerate(rows):
        row = t.rows[r]
        if r == 0:
            trpr = row._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            trpr.append(header)
        for c, width in enumerate(widths):
            cell = row.cells[c]
            cell.width = Cm(width)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            inline(p, cells[c] if c < len(cells) else "", size=9, bold=r == 0)
            if r == 0:
                shade(cell._tc, "E3EBF6")
    borders = t._tbl.tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        t._tbl.tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "C9D2DE")
        borders.append(el)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def image(doc: Document, caption: str, path: Path) -> None:
    width_px, height_px = Image.open(path).size
    width = TEXT_WIDTH_CM
    if height_px / width_px * width > 21:  # высокий снимок не должен занять больше страницы
        width = 21 * width_px / height_px
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(str(path), width=Cm(width))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(caption)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def list_item(doc: Document, text: str, marker: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
    p.add_run(marker + "\t")
    inline(p, text)


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").split("\n")
    doc = Document()
    setup(doc)

    title = lines[0].lstrip("# ").strip()
    intro: list[str] = []
    i = 1
    while not lines[i].startswith("## "):
        if lines[i].strip():
            intro.append(lines[i].strip())
        i += 1
    title_page(doc, title, intro)

    while i < len(lines):
        line = lines[i]
        if line.startswith("## "):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("```"):
            block = []
            i += 1
            while not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            code_block(doc, block)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table(doc, rows)
            continue
        elif m := re.match(r"!\[(.*)\]\((.*)\)", line):
            image(doc, m.group(1), DOCS / m.group(2))
        elif line.startswith("- "):
            list_item(doc, line[2:], "•")
        elif m := re.match(r"(\d+)\. (.*)", line):
            list_item(doc, m.group(2), m.group(1) + ".")
        elif line.strip():
            inline(doc.add_paragraph(), line.strip())
        i += 1

    doc.core_properties.title = title
    doc.core_properties.author = "Команда unicorn"
    doc.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    build()
