"""Сохраненная версия воспроизводится как была (ТЗ, п. 3.1.5): отчет из кабинета и по ссылке
собирается из сохраненного результата, а не пересчетом на текущих нормативах, и в нем указана версия."""

from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from app.db import get_session
from app.main import app
from app.reports import pdf
from app.services import projects
from app.storage.models import User
from tests.api.test_admin_norms import LEGAL, admin  # noqa: F401  фикстура админа на базе теста
from tests.api.test_projects import IVAN, MASHA, new_project, sign_in
from tests.api.test_reports import _years_total

INSURANCE = "economics.payroll.insurance_rate"
STATE = {
    "facilityId": "warehouse",
    "taskIds": ["pallet_transport"],
    "picks": {"pallet_transport": [{"robotId": "ronavi-h1500", "share": 1}]},
    "overrides": {},
}


def preview(client) -> dict:
    request = projects.request_from(STATE).model_dump(mode="json")
    response = client.post("/api/calculations/preview", json=request)
    assert response.status_code == 200, response.text
    return response.json()


def purchase(result: dict) -> dict:
    return next(s for s in result["scenarios"] if s["id"] == "purchase")


def report(client, project_id: int, kind: str, number: int = 1):
    return client.get(f"/api/projects/{project_id}/versions/{number}/report/{kind}")


def edit_norm(client, value: float) -> None:
    client.post("/api/auth/demo", json={"role": "admin"})
    assert client.put(f"/api/admin/norms/{INSURANCE}", json={"value": value, **LEGAL}).status_code == 200
    sign_in(client, IVAN)


def test_saved_version_report_keeps_old_numbers_after_norm_change(admin):  # noqa: F811
    sign_in(admin, IVAN)
    saved = preview(admin)
    project = new_project(admin, state=STATE, result=saved)
    # на экране версии отпечаток словами, без служебных номеров файлов
    assert project["current"]["data_note"].startswith("решений в каталоге ")
    assert "файлы" not in project["current"]["data_note"]

    edit_norm(admin, 0.35)
    fresh = preview(admin)
    # норматив поменял расчет: свежие цифры другие, сохраненные остаются как были
    assert purchase(fresh)["tco_rub"] != purchase(saved)["tco_rub"]
    view = admin.get(f"/api/projects/{project['id']}").json()
    assert view["current"]["same_data"] is False

    response = report(admin, project["id"], "xlsx")
    assert response.status_code == 200, response.text
    assert "-v1-" in response.headers["content-disposition"]
    book = load_workbook(BytesIO(response.content))
    assert abs(_years_total(book, "Покупка") - purchase(saved)["tco_rub"]) < 0.01
    sources = {row[7]: row[2] for row in book["Источники"].iter_rows(min_row=6, values_only=True)}
    assert sources[INSURANCE] == 0.3
    stamp = book["Итог"]["A2"].value
    assert "версия 1" in stamp and f"модель {saved['model_version']}" in stamp
    assert view["current"]["data_version"] not in stamp
    assert "Нормативы обновились после расчета" in book["Итог"]["A4"].value
    # разделы, которые считаются заново, на других нормативах в отчет не идут
    assert "Что будет, если" not in book.sheetnames

    pdf = report(admin, project["id"], "pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


def test_report_of_a_version_on_the_same_data_is_full(admin):  # noqa: F811
    sign_in(admin, IVAN)
    saved = preview(admin)
    project = new_project(admin, state=STATE, result=saved)
    book = load_workbook(BytesIO(report(admin, project["id"], "xlsx").content))
    assert "Что будет, если" in book.sheetnames
    assert abs(_years_total(book, "Покупка") - purchase(saved)["tco_rub"]) < 0.01


def test_recalculate_makes_a_new_version_and_keeps_the_old_one(admin):  # noqa: F811
    sign_in(admin, IVAN)
    saved = preview(admin)
    project = new_project(admin, state=STATE, result=saved)
    edit_norm(admin, 0.35)

    response = admin.post(f"/api/projects/{project['id']}/versions/1/recalculate")
    assert response.status_code == 201, response.text
    second = response.json()
    assert second["number"] == 2 and second["same_data"] is True
    assert "правка" in second["data_note"]
    assert second["state"] == STATE
    assert purchase(second["result"])["tco_rub"] == purchase(preview(admin))["tco_rub"]

    first = admin.get(f"/api/projects/{project['id']}/versions/1").json()
    assert first["result"] == saved and first["same_data"] is False


def test_old_version_without_result_asks_to_recalculate(admin):  # noqa: F811
    sign_in(admin, IVAN)
    project = new_project(admin, state=STATE)
    response = report(admin, project["id"], "pdf")
    assert response.status_code == 409
    assert "Пересчитайте" in response.json()["detail"]

    assert admin.post(f"/api/projects/{project['id']}/versions/1/recalculate").status_code == 201
    assert report(admin, project["id"], "xlsx", number=2).status_code == 200
    # старая версия так и осталась без цифр
    assert admin.get(f"/api/projects/{project['id']}/versions/1").json()["result"] is None


def test_version_saved_in_an_old_shape_is_explained(db_client):
    sign_in(db_client, IVAN)
    project = new_project(db_client, state=STATE, result={"scenarios": {"purchase": {"payback_years": 2.4}}})
    response = report(db_client, project["id"], "xlsx")
    assert response.status_code == 409
    assert "старом виде" in response.json()["detail"]


def test_someone_elses_version_has_no_report_and_no_recalculation(db_client):
    sign_in(db_client, IVAN)
    project = new_project(db_client, state=STATE, result=preview(db_client))
    sign_in(db_client, MASHA)
    assert report(db_client, project["id"], "pdf").status_code == 404
    assert report(db_client, project["id"], "xlsx").status_code == 404
    assert db_client.post(f"/api/projects/{project['id']}/versions/1/recalculate").status_code == 404
    db_client.post("/api/auth/logout")
    assert report(db_client, project["id"], "pdf").status_code == 401
    assert db_client.post(f"/api/projects/{project['id']}/versions/1/recalculate").status_code == 401
    sign_in(db_client, IVAN)
    assert report(db_client, project["id"], "pdf", number=7).status_code == 404
    assert report(db_client, project["id"], "docx").status_code == 422


def test_shared_link_report_uses_saved_numbers_until_revoked(db_client):
    sign_in(db_client, IVAN)
    saved = preview(db_client)
    project = new_project(db_client, state=STATE, result=saved)
    token = db_client.post(f"/api/projects/{project['id']}/shares").json()["token"]
    db_client.post("/api/auth/logout")

    response = db_client.get(f"/api/share/{token}/report/xlsx")
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    assert abs(_years_total(book, "Покупка") - purchase(saved)["tco_rub"]) < 0.01

    sign_in(db_client, IVAN)
    db_client.delete(f"/api/projects/{project['id']}/shares/{token}")
    db_client.post("/api/auth/logout")
    assert db_client.get(f"/api/share/{token}/report/pdf").status_code == 404


def test_pdf_first_page_says_catalog_and_norms_in_words_and_keeps_fingerprint_for_the_end(db_client):
    sign_in(db_client, IVAN)
    project = new_project(db_client, state=STATE, result=preview(db_client))
    fingerprint = project["current"]["data_version"]
    session = next(app.dependency_overrides[get_session]())
    user = session.scalar(select(User).where(User.email == IVAN["email"]))
    report = projects.version_report(session, user, project["id"], 1)

    pdf._register()
    first_page = " ".join(part.text for part in pdf._version(report, pdf._styles()) if hasattr(part, "text"))
    assert "Каталог и нормативы: решений в каталоге" in first_page
    assert fingerprint not in first_page and "файлы" not in first_page
    assert fingerprint not in report.stamp
    assert report.data_version == fingerprint


def test_readable_note_drops_file_fingerprint_also_in_old_versions():
    from app.services.versions import readable_note

    stored = (
        "модель 0.2, файлы 49326edaf0a3; каталог: решений 190, правка 6; нормативы: как в файлах, правок в админке нет"
    )
    assert readable_note(stored) == (
        "решений в каталоге 190, правка каталога 6; нормативы: как в файлах, правок в админке нет"
    )
    assert readable_note("") == ""
