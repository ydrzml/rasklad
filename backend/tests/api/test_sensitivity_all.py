"""Что будет, если по всем числам: что берем, как сдвигаем, как сортируем (app/services/sensitivity.py)."""

import pytest
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)
PALLETS = "facilities.warehouse.operations.pallet_transport"


def post(url: str, **body) -> dict:
    response = client.post(url, json=body)
    assert response.status_code == 200, response.text
    return response.json()


@pytest.fixture(scope="module")
def full() -> dict:
    return post("/api/calculations/sensitivity/all")


def by_id(result: dict, path: str) -> dict:
    return next(p for p in result["params"] if p["id"] == path)


def swing(param: dict, scenario: str = "purchase") -> float:
    return next(s["rub"] for s in param["swings"] if s["scenario_id"] == scenario)


def test_sorted_by_swing_and_people_first(full):
    swings = [max(s["rub"] for s in p["swings"]) for p in full["params"]]
    assert swings == sorted(swings, reverse=True)
    # роботы забирают работу 25 водителей погрузчиков: какую долю их работы, их число и оклад двигают итог
    # сильнее всего. Люди двигаются целыми, но 25 -20% это ровно 20, +20% ровно 30: число людей сдвигается
    # на те же 20%, что доля работы, и итог у них одинаковый, порядок между ними не проверяем. Оклад слабее:
    # по нему же платим присмотру (оператор это переученный водитель), и часть сдвига уходит туда
    top = [p["id"] for p in full["params"][:3]]
    labor = f"{PALLETS}.labor_share"
    headcount = "facilities.warehouse.staff.forklift-operators.headcount"
    assert set(top[:2]) == {labor, headcount}
    assert swing(by_id(full, labor)) == pytest.approx(swing(by_id(full, headcount)))
    assert top[2] == "facilities.warehouse.staff.forklift-operators.salary_month"


def test_rules_and_checks_are_not_moved(full):
    moved = {p["id"] for p in full["params"]}
    assert not any(".constraints." in path or path.startswith("engine.simulation") for path in moved)
    assert "economics.horizon_years" not in moved and "economics.discount_rate" not in moved
    assert "facilities.warehouse.schedule.shifts" not in moved  # 2 смены плюс 10% не бывает
    # нули в процентах не сдвинуть: их перечисляем отдельно, а не молча выбрасываем
    assert "Надбавка оператору роботов" in full["zeros"]
    # оклад отборщика в расчете перевозки паллет ни на что не влияет
    assert "Оклад в месяц, отборщик (комплектовщик)" in full["flat"]


def test_same_numbers_as_main_table(full):
    main = post("/api/calculations/sensitivity")
    price = next(p for p in main["params"] if p["id"] == "robot_price")
    assert by_id(full, "robots.ronavi-h1500.price_rub")["cells"] == price["cells"]


def test_share_does_not_go_above_one():
    # доля 95%: +10% и +20% дают одно и то же, 100%
    param = by_id(post("/api/calculations/sensitivity/all"), f"{PALLETS}.automatable_share")
    assert param["values"] == pytest.approx([0.76, 0.855, 0.95, 1.0, 1.0])
    plus10, plus20 = ({k: v for k, v in cell.items() if k != "delta"} for cell in param["cells"][3:])
    assert plus10 == plus20


def test_people_move_in_whole_persons():
    # 25 человек -10% это 23 (22,5 вверх), а не 22,5: +10% это 28 (27,5 вверх)
    body = {"overrides": {}}
    base = post("/api/calculations/preview", **body)["scenarios"][0]["tco_rub"]
    param = by_id(
        post("/api/calculations/sensitivity/all", **body), "facilities.warehouse.staff.forklift-operators.headcount"
    )
    minus, plus = param["cells"][1]["baseline_tco_rub"], param["cells"][3]["baseline_tco_rub"]
    # без роботов стоимость пропорциональна числу людей роли плюс остальной штат
    assert (base - minus) / (plus - base) == pytest.approx(2 / 3, rel=1e-6)


def test_fleet_from_shift_follows_the_formula():
    """Парк из прогона меняется во столько раз, во сколько по формуле меняется потребность.
    Без плана маршрут 75 м, спрос с резервом 1900 / 22 * 1,5 * 1,15 = 148,98 в час. Скорость +20%: по формуле
    148,98 / (3600 / (2 * 75 / 1,2 + 20) * 0,75) = 8,0 робота вместо 148,98 / 15,88 = 9,38, прогон давал 9,
    9 * 8,0 / 9,38 = 7,68, вверх 8. Скорость -20%: 11,45 робота, 9 * 11,45 / 9,38 = 10,99, вверх 11."""
    result = post("/api/calculations/sensitivity/all", use_simulation=True)
    speed = by_id(result, "robots.ronavi-h1500.avg_speed_m_s")
    assert [cell["fleet"] for cell in speed["cells"]] == [11, 10, 9, 9, 8]
    volume = by_id(result, f"{PALLETS}.volume_per_day")
    assert [cell["fleet"] for cell in volume["cells"]] == [8, 9, 9, 10, 11]


def test_shift_sees_its_own_numbers():
    """Загрузку робота и поправку маршрута прогон не читает: с прогоном они парк не двигают, их нет.
    Роботов в строю читает только прогон: 9 роботов при 90% в строю и 72% это 9 * 0,9 / 0,72 = 11,25, 12."""
    shift = {p["id"]: p for p in post("/api/calculations/sensitivity/all", use_simulation=True)["params"]}
    assert "robots.ronavi-h1500.utilization" not in shift
    assert "engine.geometry.route_factor" not in shift
    ready = shift["engine.simulation.availability"]
    assert ready["name"] == "Роботов в строю"
    assert [cell["fleet"] for cell in ready["cells"]] == [12, 10, 9, 9, 9]
    formula = {p["id"] for p in post("/api/calculations/sensitivity/all")["params"]}
    assert "robots.ronavi-h1500.utilization" in formula and "engine.simulation.availability" not in formula


def test_two_robots_are_told_apart():
    tasks = [
        {"operation_id": "pallet_transport", "robot_id": "ronavi-h1500"},
        {"operation_id": "cleaning", "robot_id": "clinbotics-400-pro"},
    ]
    names = [p["name"] for p in post("/api/calculations/sensitivity/all", tasks=tasks)["params"]]
    assert "Цена робота, H1500" in names
    assert "Цена робота, Клинботикс 400 PRO" in names


def test_excel_holds_the_full_list():
    from io import BytesIO

    from openpyxl import load_workbook

    response = client.post("/api/reports/xlsx", json={})
    sheet = load_workbook(BytesIO(response.content))["Что будет, если"]
    column = [row[0] for row in sheet.iter_rows(values_only=True)]
    assert "Все числа по силе влияния" in column
    start = column.index("Все числа по силе влияния")
    assert column[start + 3] == "Доля работы людей, которую забирает робот, перевозка паллет между зонами"
