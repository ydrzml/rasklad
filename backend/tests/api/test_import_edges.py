"""Импорт данных объекта: верный файл, значения вне диапазона, пустые и битые файлы."""

import zipfile
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app
from app.services import data_import

client = TestClient(app)
AREA = "facilities.warehouse.active_area_m2"
SHIFTS = "facilities.warehouse.schedule.shifts"
SHIFT_HOURS = "facilities.warehouse.schedule.shift_hours"


def template() -> Workbook:
    response = client.post("/api/import/template", json={})
    assert response.status_code == 200, response.text
    return load_workbook(BytesIO(response.content))


def upload(content: bytes, name: str, **form):
    return client.post("/api/import", files={"file": (name, content)}, data={"facility_id": "warehouse", **form})


def saved(book: Workbook) -> bytes:
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def field_range(book: Workbook, path: str) -> tuple[float, float]:
    for row in book["Объект"].iter_rows(min_row=2, values_only=True):
        if row[0] == path:
            return float(row[4]), float(row[5])
    raise AssertionError(f"нет строки {path} в шаблоне")


def set_value(book: Workbook, path: str, value) -> None:
    for row in book["Объект"].iter_rows(min_row=2):
        if row[0].value == path:
            row[2].value = value
            return
    raise AssertionError(f"нет строки {path} в шаблоне")


def by_field(result: dict) -> dict[str, dict]:
    return {row["field"]: row for row in result["rows"]}


# --- Верный файл ----------------------------------------------------------------------


def test_untouched_template_is_all_ok_and_applies_every_row():
    book = template()
    result = upload(saved(book), "шаблон.xlsx").json()
    assert result["errors"] == 0 and result["warnings"] == 0
    assert all(row["status"] == "ok" for row in result["rows"])
    object_rows = book["Объект"].max_row - 1
    staff_rows = book["Штат"].max_row - 1
    assert len(result["overrides"]) == object_rows
    assert len(result["staff"]) == staff_rows
    assert result["applied"] == object_rows + staff_rows


def test_template_carries_the_range_for_every_parameter():
    book = template()
    for row in book["Объект"].iter_rows(min_row=2, values_only=True):
        path, _label, value, unit, low, high = row[:6]
        assert path and unit, row
        assert low is not None and high is not None and float(low) <= float(high), row
        assert float(low) <= float(value) <= float(high), f"значение по умолчанию вне диапазона: {row}"


def test_csv_with_utf8_bom_and_comma_separator():
    text = "\ufeffПараметр,Значение\nПлощадь зоны работы роботов,12000\n"
    result = upload(text.encode("utf-8"), "данные.csv").json()
    assert result["overrides"] == {AREA: 12000}
    assert result["errors"] == 0


def test_old_field_name_still_imports():
    # поле переименовали, а файлы по прежнему шаблону у людей остались
    text = "Параметр;Значение\nПлощадь активной зоны;12000\n"
    result = upload(text.encode("utf-8"), "данные.csv").json()
    assert result["overrides"] == {AREA: 12000}
    assert by_field(result)["Площадь зоны работы роботов"]["status"] == "ok"


def test_extension_case_does_not_matter():
    text = "Параметр;Значение\nПлощадь зоны работы роботов;12000\n"
    assert upload(text.encode("utf-8"), "ДАННЫЕ.CSV").status_code == 200
    assert upload(saved(template()), "ШАБЛОН.XLSX").status_code == 200


# --- Значения вне диапазона и прочие ошибки строк -------------------------------------


def test_value_out_of_range_is_applied_with_a_warning():
    book = template()
    _low, high = field_range(book, AREA)
    set_value(book, AREA, high * 10)
    result = upload(saved(book), "склад.xlsx").json()
    row = by_field(result)["Площадь зоны работы роботов"]
    assert row["status"] == "warn" and "Вне того" in row["message"]
    assert result["overrides"][AREA] == high * 10
    assert result["warnings"] == 1 and result["errors"] == 0


def test_value_below_range_is_a_warning_too():
    book = template()
    low, _high = field_range(book, AREA)
    if low > 0:
        set_value(book, AREA, low / 10)
        result = upload(saved(book), "склад.xlsx").json()
        assert by_field(result)["Площадь зоны работы роботов"]["status"] == "warn"
        assert result["overrides"][AREA] == low / 10


def test_negative_text_and_wrong_unit_are_errors_and_not_applied():
    book = template()
    sheet = book["Объект"]
    for row in sheet.iter_rows(min_row=2):
        if row[0].value == AREA:
            row[2].value = -5
        if row[0].value == SHIFT_HOURS:
            row[2].value = "восемь"
    sheet.append(["", "Площадь зоны работы роботов", 500, "га"])
    result = upload(saved(book), "склад.xlsx").json()
    messages = [(row["field"], row["status"], row["message"]) for row in result["rows"]]
    assert ("Площадь зоны работы роботов", "error", "Меньше нуля быть не может") in messages
    assert any(
        field == "Продолжительность смены" and status == "error" and "Не число" in message
        for field, status, message in messages
    )
    assert any(field == "Площадь зоны работы роботов" and "га" in message for field, status, message in messages)
    assert AREA not in result["overrides"] and SHIFT_HOURS not in result["overrides"]
    assert result["errors"] == 3


def test_repeated_parameter_takes_the_last_value_with_a_warning():
    text = "Параметр;Значение\nПлощадь зоны работы роботов;12000\nПлощадь зоны работы роботов;13000\n"
    result = upload(text.encode("utf-8"), "данные.csv").json()
    assert result["overrides"][AREA] == 13000
    # строки файла, без предупреждений о главных полях, которых в файле нет
    statuses = [row["status"] for row in result["rows"] if row["row"]]
    assert statuses == ["ok", "warn"]


def test_schedule_over_24_hours_is_refused_as_a_whole():
    # 3 смены по 11 часов это 33 часа в сутках: обе строки ошибкой, ни одна не подставляется
    text = f"Код;Значение\n{SHIFTS};3\n{SHIFT_HOURS};11\n"
    result = upload(text.encode("utf-8"), "данные.csv").json()
    assert SHIFTS not in result["overrides"] and SHIFT_HOURS not in result["overrides"]
    assert all(row["status"] == "error" and "24" in row["message"] for row in result["rows"] if row["row"])


def test_staff_rows_with_bad_numbers_or_flags_are_skipped():
    book = template()
    people = book["Штат"]
    role = people.cell(row=2, column=1).value
    people.append([role, -1, 0, 100000, "нет"])
    people.append([role, 5, 3, 100000, "может быть"])
    people.append([role, 5, 7, None, "нет"])
    result = upload(saved(book), "штат.xlsx").json()
    tail = result["rows"][-3:]
    assert tail[0]["status"] == "error" and "не меньше нуля" in tail[0]["message"]
    assert tail[1]["status"] == "error" and "да" in tail[1]["message"]
    assert tail[2]["status"] == "warn" and "больше, чем мест" in tail[2]["message"]
    assert "рыночный" in tail[2]["message"]
    assert len(result["staff"]) == people.max_row - 1 - 2


# --- Пустые и битые файлы -------------------------------------------------------------


def test_empty_csv_and_header_only_csv_apply_nothing():
    for content in (b"", "Параметр;Значение\n".encode(), b"   \n\n"):
        response = upload(content, "пусто.csv")
        assert response.status_code == 200, response.text
        result = response.json()
        assert result["applied"] == 0 and result["overrides"] == {} and result["staff"] is None


def test_csv_without_our_columns_explains_what_to_do():
    result = upload(b"a;b\n1;2\n", "чужой.csv").json()
    assert result["applied"] == 0 and result["errors"] == 1
    assert "шаблон" in result["rows"][0]["message"].lower()


def test_empty_workbook_applies_nothing():
    response = upload(saved(Workbook()), "пустая.xlsx")
    assert response.status_code == 200, response.text
    assert response.json()["applied"] == 0


def test_corrupt_xlsx_is_refused_with_a_reason():
    for content in (b"", b"not a zip at all", saved(template())[:200]):
        response = upload(content, "битый.xlsx")
        assert response.status_code == 415, response.text
        assert "поврежден" in response.json()["detail"]


def test_zip_that_is_not_a_workbook_is_refused():
    out = BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("readme.txt", "hello")
    assert upload(out.getvalue(), "архив.xlsx").status_code == 415


def test_binary_garbage_named_csv_does_not_crash():
    response = upload(bytes(range(256)) * 4, "мусор.csv")
    assert response.status_code in (200, 415), response.text
    if response.status_code == 200:
        assert response.json()["overrides"] == {}


def test_file_over_the_limit_is_refused():
    too_big = b"x" * (data_import.MAX_BYTES + 1)
    assert upload(too_big, "большой.csv").status_code == 413
    assert upload(too_big, "большой.xlsx").status_code == 413


def test_unknown_extension_and_no_name_are_refused():
    header = "Параметр;Значение\n".encode()
    assert upload(header, "данные.txt").status_code == 415
    assert upload(header, "").status_code in (415, 422)


def test_unknown_facility_is_not_found():
    text = "Параметр;Значение\nПлощадь зоны работы роботов;12000\n"
    assert upload(text.encode("utf-8"), "данные.csv", facility_id="космодром").status_code == 404


# --- Пустые и пропущенные обязательные поля -------------------------------------------


def test_empty_required_field_is_an_error_with_the_default_named():
    text = "Параметр;Значение\nПлощадь зоны работы роботов;\nРабочих дней в году;\n"
    result = upload(text.encode("utf-8"), "данные.csv").json()
    area, days = result["rows"][0], result["rows"][1]
    assert (
        area["status"] == "error"
        and "Обязательное поле пустое" in area["message"]
        and "10\u00a0000\u00a0м2" in area["message"]
    )
    assert days["status"] == "warn" and "по умолчанию: 365" in days["message"]
    assert AREA not in result["overrides"]
    # пустые поля ничего не подставляют и в "подставим с замечанием" не считаются
    assert result["applied"] == 0 and result["warnings"] == 0


def test_required_field_missing_from_file_is_named():
    text = "Параметр;Значение\nРабочих дней в году;300\n"
    result = upload(text.encode("utf-8"), "данные.csv").json()
    missing = [row for row in result["rows"] if not row["row"]]
    assert {row["field"] for row in missing} >= {
        "Площадь зоны работы роботов",
        "Смен в сутки",
        "Продолжительность смены",
    }
    assert all(row["status"] == "warn" and "нет в файле" in row["message"] for row in missing)
