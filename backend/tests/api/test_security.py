"""Проверка безопасности: доступ ко всем точкам API, сессия и пароль, Excel без чужих формул.

Список точек берем из описания API самого приложения. Новая точка должна либо требовать вход,
либо быть записана в PUBLIC: иначе тест упадет, и про нее придется подумать."""

import re
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.auth import service
from app.db import get_session
from app.main import app
from app.schemas.plan import MAX_PLAN_ITEMS
from app.schemas.staff import MAX_STAFF_LINES
from app.services import projects
from app.settings import settings
from tests.api.test_projects import IVAN, MASHA, RESULT, STATE, new_project, sign_in

# Что открыто без входа: справочники, расчет, отчеты по введенным цифрам, вход и ссылка "Поделиться"
PUBLIC = {
    ("GET", "/api/health"),
    ("GET", "/api/auth/me"),
    ("GET", "/api/auth/demo"),
    ("POST", "/api/auth/demo"),
    ("POST", "/api/auth/login"),
    ("POST", "/api/auth/register"),
    ("POST", "/api/auth/logout"),
    ("GET", "/api/catalog/facilities"),
    ("GET", "/api/catalog/solutions"),
    ("GET", "/api/catalog/parameters"),
    ("GET", "/api/catalog/robot-parameters"),
    ("GET", "/api/catalog/operations"),
    ("POST", "/api/catalog/facility-solutions"),
    ("GET", "/api/catalog/photos/{solution_id}"),
    ("GET", "/api/staff"),
    ("POST", "/api/staff/check"),
    ("POST", "/api/staff/form"),
    ("POST", "/api/budget/fit"),
    ("POST", "/api/import/template"),
    ("POST", "/api/import"),
    ("GET", "/api/plan/templates"),
    ("GET", "/api/plan/rack-types"),
    ("POST", "/api/plan/generate"),
    ("POST", "/api/plan/measure"),
    ("POST", "/api/plan/serpentine"),
    ("POST", "/api/calculations/preview"),
    ("POST", "/api/calculations/sensitivity"),
    ("POST", "/api/calculations/sensitivity/all"),
    ("POST", "/api/simulation/runs"),
    ("POST", "/api/simulation/log.csv"),
    ("POST", "/api/reports/pdf"),
    ("POST", "/api/reports/xlsx"),
    ("POST", "/api/readiness"),
    ("GET", "/api/share/{token}"),
}

# Тела, без которых сервер ответил бы 422 раньше, чем проверил вход
BODIES = {
    "/api/auth/me/password": {"password": "whatever-1", "new_password": "whatever-2"},
    "/api/auth/me/email": {"email": "new@example.com", "password": "whatever-1"},
}


def operations() -> list[tuple[str, str]]:
    return [(method.upper(), path) for path, methods in app.openapi()["paths"].items() for method in methods]


def url(path: str, value: str = "1") -> str:
    return re.sub(r"\{[^}]+\}", value, path)


# --- Гость ------------------------------------------------------------------------------


def test_public_list_is_not_stale():
    # Если точку из списка удалили или переименовали, список надо поправить
    assert PUBLIC <= set(operations())


def test_guest_gets_401_on_every_closed_point(db_client):
    closed = [op for op in operations() if op not in PUBLIC]
    assert len(closed) >= 40
    for method, path in closed:
        response = db_client.request(method, url(path), json=BODIES.get(path))
        assert response.status_code == 401, (method, path, response.status_code)


def test_user_gets_403_on_every_admin_point(db_client):
    sign_in(db_client, IVAN)
    admin = [op for op in operations() if op[1].startswith("/api/admin")]
    assert len(admin) >= 20
    for method, path in admin:
        response = db_client.request(method, url(path))
        assert response.status_code == 403, (method, path, response.status_code)


# --- Чужие проекты ----------------------------------------------------------------------


def ivans_things(client) -> dict:
    sign_in(client, IVAN)
    project = new_project(client, name="Склад Ивана", state=STATE, result=RESULT)
    base = f"/api/projects/{project['id']}"
    file_id = client.post(f"{base}/files", files={"file": ("штат.csv", b"role;count\n", "text/csv")}).json()["id"]
    token = client.post(f"{base}/shares").json()["token"]
    folder = client.post("/api/projects/folders", json={"name": "Папка Ивана", "projects": [project["id"]]}).json()
    return {"project": project["id"], "file": file_id, "token": token, "folder": folder["id"]}


def test_stranger_can_neither_see_nor_change_anything(db_client):
    ivan = ivans_things(db_client)
    sign_in(db_client, MASHA)
    own = new_project(db_client, name="Склад Маши", state=STATE, result=RESULT)
    base = f"/api/projects/{ivan['project']}"
    calls = [
        ("GET", base, None),
        ("PATCH", base, {"name": "взлом"}),
        ("DELETE", base, None),
        ("POST", f"{base}/copy", {}),
        ("POST", f"{base}/versions", {"state": {}}),
        ("GET", f"{base}/versions/1", None),
        ("POST", f"{base}/shares", {"version": 1}),
        ("DELETE", f"{base}/shares/{ivan['token']}", None),
        ("GET", f"{base}/files/{ivan['file']}", None),
        ("DELETE", f"{base}/files/{ivan['file']}", None),
        ("PATCH", f"/api/projects/folders/{ivan['folder']}", {"name": "взлом"}),
        ("DELETE", f"/api/projects/folders/{ivan['folder']}", None),
        # свой проект в чужую папку и чужой проект в свою папку
        ("PATCH", f"/api/projects/{own['id']}", {"folder_id": ivan["folder"]}),
        ("POST", "/api/projects/folders", {"name": "Моя", "projects": [ivan["project"]]}),
        ("POST", "/api/projects/compare", {"items": [{"project_id": own["id"]}, {"project_id": ivan["project"]}]}),
    ]
    for method, path, body in calls:
        response = db_client.request(method, path, json=body)
        assert response.status_code == 404, (method, path, response.status_code)
    upload = db_client.post(f"{base}/files", files={"file": ("x.csv", b"a;b\n", "text/csv")})
    assert upload.status_code == 404

    # у Маши ничего чужого не появилось, в том числе папки от неудачной попытки
    assert [p["name"] for p in db_client.get("/api/projects").json()] == ["Склад Маши"]
    assert db_client.get("/api/projects/folders").json() == []

    # у Ивана все на месте
    sign_in(db_client, IVAN)
    project = db_client.get(base).json()
    assert project["name"] == "Склад Ивана"
    assert project["versions"] == 1
    assert len(project["files"]) == 1
    assert [f["name"] for f in db_client.get("/api/projects/folders").json()] == ["Папка Ивана"]
    assert db_client.get(f"/api/share/{ivan['token']}").status_code == 200


def test_shared_link_shows_no_owner_data(db_client):
    ivan = ivans_things(db_client)
    db_client.patch("/api/auth/me", json={"name": "Иван Петров", "company": "ООО Склад"})
    db_client.post("/api/auth/logout")
    text = db_client.get(f"/api/share/{ivan['token']}").text
    for secret in (IVAN["email"], "Иван Петров", "ООО Склад", "owner", "password"):
        assert secret not in text


# --- Сессия и пароль --------------------------------------------------------------------


def test_session_cookie_is_closed_for_scripts_and_other_sites(db_client, monkeypatch):
    cookie = db_client.post("/api/auth/register", json=IVAN).headers["set-cookie"].lower()
    assert "httponly" in cookie
    assert "samesite=lax" in cookie
    assert "secure" not in cookie

    monkeypatch.setattr(settings, "cookie_secure", True)
    cookie = db_client.post("/api/auth/login", json=IVAN).headers["set-cookie"].lower()
    assert "; secure" in cookie


def test_password_change_signs_out_other_devices(db_client):
    sign_in(db_client, IVAN)
    laptop = TestClient(app)
    laptop.post("/api/auth/login", json=IVAN)
    assert laptop.get("/api/auth/me").json()["role"] == "user"

    new = "brand-new-secret"
    changed = db_client.post("/api/auth/me/password", json={"password": IVAN["password"], "new_password": new})
    assert changed.status_code == 200
    # здесь вход остался, а украденная или забытая на другом компьютере cookie больше не действует
    assert db_client.get("/api/auth/me").json()["role"] == "user"
    assert laptop.get("/api/auth/me").json()["role"] == "guest"
    assert laptop.get("/api/projects").status_code == 401
    assert db_client.post("/api/auth/login", json={**IVAN, "password": new}).status_code == 200


def test_email_and_password_change_need_the_current_password_and_stop_guessing(db_client):
    sign_in(db_client, IVAN)
    wrong = {"password": "not-my-password", "new_password": "another-secret"}
    for _ in range(10):
        assert db_client.post("/api/auth/me/password", json=wrong).status_code == 401
    assert db_client.post("/api/auth/me/password", json=wrong).status_code == 429
    for _ in range(10):
        response = db_client.post("/api/auth/me/email", json={"email": "x@example.com", "password": "nope-nope"})
        assert response.status_code == 401
    assert db_client.post("/api/auth/me/email", json={"email": "x@example.com", "password": "nope"}).status_code == 429


def test_errors_do_not_show_code_or_database(db_client):
    client = TestClient(app, raise_server_exceptions=False)
    for response in (
        client.post("/api/auth/login", json={"email": "x@example.com", "password": "' or 1=1 --"}),
        client.get("/api/projects/1%27%20or%201=1"),
        client.post("/api/reports/pdf", json={"robot_id": "../../etc/passwd"}),
    ):
        text = response.text.lower()
        for leak in ("traceback", "sqlalchemy", "psycopg", "sqlite", 'file "'):
            assert leak not in text, (response.request.url, leak)


# --- Отчеты -----------------------------------------------------------------------------


def test_excel_does_not_run_formulas_from_user_text(db_client):
    staff = [{"id": "x", "role": '=HYPERLINK("http://evil.example","открыть")', "headcount": 5, "filled": 5}]
    staff[0]["salary_month"] = 80000
    response = db_client.post("/api/reports/xlsx", json={"staff": staff})
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    cells = [cell for sheet in book for row in sheet.iter_rows() for cell in row]
    evil = [cell for cell in cells if isinstance(cell.value, str) and "evil" in cell.value]
    assert evil
    assert all(cell.data_type == "s" for cell in evil)
    # свои формулы остались формулами
    assert any(cell.data_type == "f" for cell in cells)


def test_import_template_does_not_run_formulas_from_user_text(db_client):
    staff = [{"id": "x", "role": "=cmd|' /C calc'!A0", "headcount": 1, "filled": 1, "salary_month": 1}]
    response = db_client.post("/api/import/template", json={"staff": staff})
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    evil = [cell for sheet in book for row in sheet.iter_rows() for cell in row if "cmd" in str(cell.value)]
    assert evil
    assert all(cell.data_type == "s" for cell in evil)


# --- Потолки и пароль -------------------------------------------------------------------


def test_one_account_cannot_fill_the_server(db_client, monkeypatch):
    monkeypatch.setattr(projects, "MAX_PROJECTS", 2)
    monkeypatch.setattr(projects, "MAX_FOLDERS", 1)
    monkeypatch.setattr(projects, "MAX_VERSIONS", 3)
    monkeypatch.setattr(projects, "MAX_FILES_TOTAL_BYTES", 10)
    sign_in(db_client, IVAN)
    first = new_project(db_client, state=STATE, result=RESULT)
    base = f"/api/projects/{first['id']}"
    assert db_client.post(f"{base}/versions", json={"state": STATE}).status_code == 201
    new_project(db_client, state=STATE)

    refused = [
        db_client.post("/api/projects", json={"name": "третий"}),
        db_client.post(f"{base}/copy", json={}),
        db_client.post(f"{base}/versions", json={"state": STATE}),
        db_client.post(f"{base}/files", files={"file": ("big.csv", b"x" * 11, "text/csv")}),
    ]
    assert db_client.post("/api/projects/folders", json={"name": "одна"}).status_code == 201
    refused.append(db_client.post("/api/projects/folders", json={"name": "вторая"}))
    for response in refused:
        assert response.status_code == 422, response.request.url
        assert "предел" in response.json()["detail"]
    assert len(db_client.get("/api/projects").json()) == 2


def test_new_password_cannot_be_trivial_but_old_ones_still_log_in(db_client):
    for weak in ("12345678", "password123", "aaaaaaaa", "Qwerty123"):
        response = db_client.post("/api/auth/register", json={"email": "weak@example.com", "password": weak})
        assert response.status_code == 422, weak
        # объясняем по-русски и не возвращаем сам пароль
        assert response.json()["detail"][0]["msg"]
        assert weak not in response.text
    assert db_client.post("/api/auth/register", json=IVAN).status_code == 201
    change = {"password": IVAN["password"], "new_password": "87654321"}
    assert db_client.post("/api/auth/me/password", json=change).status_code == 422
    # аккаунт, заведенный с простым паролем раньше, входит как прежде: проверяем только новые пароли
    service.register(next(app.dependency_overrides[get_session]()), "old@example.com", "12345678")
    assert (
        db_client.post("/api/auth/login", json={"email": "old@example.com", "password": "12345678"}).status_code == 200
    )


# --- Размер ввода у открытых адресов ----------------------------------------------------


def test_open_calculations_refuse_huge_lists_before_counting(db_client):
    # Гость без входа не должен занять сервер одним запросом: штат считается чувствительностью
    # квадратично, решения в бюджете и зоны плана по одному. Отказ сразу, по-русски
    line = {"role": "picker", "headcount": 1, "filled": 1, "salary_month": 80000}
    staff = [{**line, "id": f"s{i}"} for i in range(MAX_STAFF_LINES + 1)]
    for path, body in (
        ("/api/calculations/preview", {"staff": staff}),
        ("/api/calculations/sensitivity/all", {"staff": staff}),
        ("/api/reports/xlsx", {"staff": staff}),
        ("/api/import/template", {"staff": staff}),
        ("/api/staff/check", {"staff": staff}),
        ("/api/budget/fit", {"robot_ids": ["ronavi-h1500"] * 101, "budget_rub": 1e7}),
        ("/api/plan/measure", {"width_m": 100, "length_m": 100, "items": [{"id": "x"}] * (MAX_PLAN_ITEMS + 1)}),
    ):
        response = db_client.post(path, json=body)
        assert response.status_code == 422, (path, response.status_code)
        assert "не больше" in response.text, path
