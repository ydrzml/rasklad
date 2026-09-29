"""Решение на каждую задачу: объект считается целиком, в ответе итог и разбивка по задачам.

Старые запросы с одной задачей (operation_id и robot_id) продолжают работать: так сохранены
старые проекты, так открывается страница по ссылке и так кабинет собирает отчет.
"""

import pytest
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)
TWO = [
    {"operation_id": "pallet_transport", "robot_id": "ronavi-h1500"},
    {"operation_id": "cleaning", "robot_id": "clinbotics-600"},
]


def preview(**body):
    response = client.post("/api/calculations/preview", json=body)
    assert response.status_code == 200, response.text
    return response.json()


@pytest.fixture(scope="module")
def two():
    return preview(tasks=TWO)


def test_old_request_with_one_robot_still_works():
    result = preview(operation_id="pallet_transport", robot_id="ronavi-h1500")
    assert result["feasible"]
    assert result["operation_id"] == "pallet_transport"
    assert [task["operation_id"] for task in result["tasks"]] == ["pallet_transport"]
    assert result["sizing"]["fleet"] == 10
    assert result["sizing"]["design_demand_ops_per_hour"] > 0


def test_one_task_in_the_list_is_the_same_as_the_old_request():
    old = preview(operation_id="pallet_transport", robot_id="ronavi-h1500")
    new = preview(tasks=TWO[:1])
    assert new["scenarios"] == old["scenarios"]
    assert new["sizing"] == old["sizing"]


def test_site_fleet_is_the_sum_of_tasks(two):
    """10 паллетных роботов и 1 уборщик по формуле: на объекте 11. Дежурный один на всех."""
    assert [task["sizing"]["fleet"] for task in two["tasks"]] == [10, 1]
    assert two["sizing"]["fleet"] == 11
    assert two["sizing"]["operator_posts"] == 1
    # спрос у паллет в паллетах, у уборки в квадратных метрах: складывать их нельзя
    assert two["sizing"]["design_demand_ops_per_hour"] is None
    assert two["tasks"][0]["sizing"]["design_demand_ops_per_hour"] > 0


def test_tasks_and_shared_make_the_site(two):
    """Задачи без общих затрат и общее на объект в сумме дают объект: складывает сервер, не браузер."""
    for scenario in two["scenarios"]:
        parts = [next(p for p in task["scenarios"] if p["id"] == scenario["id"]) for task in two["tasks"]]
        shared = next(p for p in two["shared"]["scenarios"] if p["id"] == scenario["id"])
        total = sum(p["investment_year0_rub"] for p in parts) + shared["investment_year0_rub"]
        assert total == pytest.approx(scenario["investment_year0_rub"])
        tco = sum(p["tco_rub"] for p in parts) + shared["tco_rub"]
        assert tco == pytest.approx(scenario["tco_rub"])


def test_shared_holds_integration_once(two):
    """Общее на объект у покупки: ПО Ronavi 1,5 млн один раз на производителя (у Вейбота ПО в цене
    робота), интеграция 5 млн, инфраструктура 3 млн, переобучение одного дежурного 5,09 * 10 000 = 50 900 ₽
    и резерв 10% от них. Руками: 9 550 900 * 1,1 = 10 505 990 ₽."""
    shared = next(p for p in two["shared"]["scenarios"] if p["id"] == "purchase")
    assert shared["capex_rub"]["integration"] == pytest.approx(5_000_000)
    assert shared["capex_rub"]["software"] == pytest.approx(1_500_000)
    assert shared["investment_year0_rub"] == pytest.approx(10_505_990, abs=1)


def test_reversed_tasks_give_the_same_site(two):
    """Итог по объекту не зависит от порядка задач: общее считается по объекту, а не по первой задаче."""
    reverse = preview(tasks=TWO[::-1])
    for one, other in zip(two["scenarios"], reverse["scenarios"], strict=True):
        assert one["tco_rub"] == pytest.approx(other["tco_rub"])
        assert one["investment_year0_rub"] == pytest.approx(other["investment_year0_rub"])
    assert reverse["sizing"]["fleet"] == two["sizing"]["fleet"]
    assert reverse["sizing"]["operator_cost_rub_year"] == pytest.approx(two["sizing"]["operator_cost_rub_year"])


def test_blocked_refund_and_leasing_on_the_whole_site():
    """Возврат Минпромторга складу недоступен: не срабатывает и говорит почему. Лизинг считается от
    вложений объекта после сложения, а разбивка по задачам сходится с вложениями до способа оплаты."""
    result = preview(tasks=TWO, subsidy_ids=["minpromtorg-robotics"], overrides={"financing.method": 2})
    purchase = next(s for s in result["scenarios"] if s["id"] == "purchase")
    refund = next(s for s in purchase["financing"]["supports"] if s["id"] == "minpromtorg-robotics")
    assert not refund["applied"] and refund["reason"].startswith("недоступна")
    assert purchase["name"] == "Покупка в лизинг"
    assert purchase["investment_year0_rub"] < purchase["capex_after_grant_rub"]
    shared = next(p for p in result["shared"]["scenarios"] if p["id"] == "purchase")
    parts = [next(p for p in task["scenarios"] if p["id"] == "purchase") for task in result["tasks"]]
    total = sum(p["investment_year0_rub"] for p in parts) + shared["investment_year0_rub"]
    assert total == pytest.approx(purchase["capex_after_grant_rub"])


def test_leasing_month_payment_and_slider_limits():
    """Панель оплаты берет платеж в месяц и границы ползунков с сервера: платеж по той же формуле,
    что график, границы из config/parameters.yaml, раздел financing."""
    result = preview(tasks=TWO, overrides={"financing.method": 2})
    plan = next(s for s in result["scenarios"] if s["id"] == "purchase")["financing"]
    i, n = plan["rate"] / 12, plan["term_months"]
    assert plan["month_first_rub"] == pytest.approx(plan["principal_rub"] * i / (1 - (1 + i) ** -n))
    assert plan["month_last_rub"] == pytest.approx(plan["month_first_rub"])
    assert plan["term_max_months"] == 84 and plan["rate_max"] == 0.35


def test_sources_hold_both_robots_and_both_tasks(two):
    paths = {source["path"] for source in two["sources"]}
    assert "robots.ronavi-h1500.price_rub" in paths
    assert "robots.clinbotics-600.price_rub" in paths
    assert "facilities.warehouse.operations.cleaning.volume_per_day" in paths
    assert not any(path.startswith("facilities.warehouse.operations.piece_picking.") for path in paths)


def test_same_task_twice_is_rejected():
    response = client.post("/api/calculations/preview", json={"tasks": [TWO[0], TWO[0]]})
    assert response.status_code == 422
    assert "дважды" in response.json()["detail"]


def test_sensitivity_moves_every_task(two):
    """Объем +20% двигает каждую задачу. Руками по формуле: паллеты 148,98 * 1,2 / 15,88 = 11,26,
    вверх 12; уборка 548,86 * 1,2 / 750 = 0,88, вверх 1. Вместе 13 вместо 11."""
    response = client.post("/api/calculations/sensitivity", json={"tasks": TWO})
    assert response.status_code == 200, response.text
    volume = next(p for p in response.json()["params"] if p["id"] == "volume")
    assert volume["cells"][2]["fleet"] == two["sizing"]["fleet"]
    assert volume["cells"][-1]["fleet"] == 13


def test_reports_are_built_for_the_whole_site():
    pdf = client.post("/api/reports/pdf", json={"tasks": TWO})
    assert pdf.status_code == 200, pdf.text
    assert pdf.content.startswith(b"%PDF")
    xlsx = client.post("/api/reports/xlsx", json={"tasks": TWO})
    assert xlsx.status_code == 200, xlsx.text

    from io import BytesIO

    from openpyxl import load_workbook

    book = load_workbook(BytesIO(xlsx.content))
    assert "По задачам" in book.sheetnames
    rows = [[cell.value for cell in row] for row in book["По задачам"].iter_rows()]
    names = [row[0] for row in rows]
    assert "Общее на объект" in names and "Объект целиком" in names


MIXED = [
    {"operation_id": "pallet_transport", "robot_id": "ronavi-h1500", "share": 0.5},
    {"operation_id": "pallet_transport", "robot_id": "moros-amr-1500", "share": 0.5},
]


def test_mixed_fleet_gives_a_part_per_solution():
    """Два решения на задачу: в ответе по части на решение с долей, парк объекта их сумма."""
    result = preview(tasks=MIXED)
    assert [(t["robot_id"], t["share"], t["sizing"]["fleet"]) for t in result["tasks"]] == [
        ("ronavi-h1500", 0.5, 5),
        ("moros-amr-1500", 0.5, 5),
    ]
    assert result["sizing"]["fleet"] == 10
    assert result["sizing"]["operator_posts"] == 1
    # источники держат оба робота
    paths = {source["path"] for source in result["sources"]}
    assert "robots.moros-amr-1500.price_rub" in paths and "robots.ronavi-h1500.price_rub" in paths


def test_mixed_fleet_shares_must_add_up():
    response = client.post(
        "/api/calculations/preview",
        json={"tasks": [{**MIXED[0], "share": 0.6}, {**MIXED[1], "share": 0.3}]},
    )
    assert response.status_code == 422
    assert "90%" in response.json()["detail"]
    same = client.post("/api/calculations/preview", json={"tasks": [MIXED[0], {**MIXED[0]}]})
    assert same.status_code == 422
    three = [
        {**MIXED[0], "share": 0.4},
        {**MIXED[1], "share": 0.3},
        {"operation_id": "pallet_transport", "robot_id": "ronavi-h2000", "share": 0.3},
    ]
    response = client.post("/api/calculations/preview", json={"tasks": three})
    assert response.status_code == 422
    assert "не больше двух" in response.json()["detail"]


def test_mixed_fleet_shift_runs_the_share():
    """Прогон смены части гоняет ее долю спроса: у половины паллет спрос вдвое меньше."""
    whole = client.post("/api/simulation/runs", json={"fleet": 5, "with_events": False}).json()
    half = client.post("/api/simulation/runs", json={"fleet": 5, "with_events": False, "share": 0.5}).json()
    assert half["design_demand"] == pytest.approx(whole["design_demand"] / 2)
    assert sum(half["demand_per_hour"]) == pytest.approx(sum(whole["demand_per_hour"]) / 2)


def test_mixed_fleet_reports_have_a_row_per_solution():
    xlsx = client.post("/api/reports/xlsx", json={"tasks": MIXED})
    assert xlsx.status_code == 200, xlsx.text
    from io import BytesIO

    from openpyxl import load_workbook

    rows = [[c.value for c in r] for r in load_workbook(BytesIO(xlsx.content))["По задачам"].iter_rows()]
    names = [row[0] for row in rows if row[0]]
    assert sum(1 for name in names if name.startswith("Перевозка паллет")) >= 2
    pdf = client.post("/api/reports/pdf", json={"tasks": MIXED})
    assert pdf.status_code == 200, pdf.text


def test_picking_sensitivity_has_station_rate_and_plan_reason():
    """У отбора в «что будет, если» есть выработка на станции; отбор на плане без станций называет причиной план."""
    response = client.post(
        "/api/calculations/sensitivity",
        json={"tasks": [{"operation_id": "piece_picking", "robot_id": "ronavi-m"}]},
    )
    assert response.status_code == 200, response.text
    ids = [param["id"] for param in response.json()["params"]]
    assert "station_rate" in ids
    pallets_only = client.post(
        "/api/calculations/sensitivity",
        json={"tasks": [{"operation_id": "pallet_transport", "robot_id": "ronavi-h1500"}]},
    ).json()
    assert "station_rate" not in [param["id"] for param in pallets_only["params"]]
