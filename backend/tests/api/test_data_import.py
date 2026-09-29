from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

client = TestClient(app)
AREA = "facilities.warehouse.active_area_m2"
SHIFTS = "facilities.warehouse.schedule.shifts"


def template(**body) -> bytes:
    response = client.post("/api/import/template", json=body)
    assert response.status_code == 200, response.text
    return response.content


def upload(content: bytes, name: str):
    return client.post("/api/import", files={"file": (name, content)}, data={"facility_id": "warehouse"})


def test_template_is_filled_with_current_values():
    book = load_workbook(BytesIO(template(overrides={AREA: 12345})))
    assert book.sheetnames == ["Объект", "Штат", "Роли"]
    rows = {row[0]: row for row in book["Объект"].iter_rows(min_row=2, values_only=True)}
    assert rows[AREA][2] == 12345 and rows[AREA][3] == "м2"
    assert book["Штат"].max_row > 1


def test_excel_round_trip_reports_every_row():
    book = load_workbook(BytesIO(template()))
    sheet = book["Объект"]
    for row in sheet.iter_rows(min_row=2):
        if row[0].value == AREA:
            row[2].value = 12000  # в диапазоне
        if row[0].value == SHIFTS:
            row[2].value = 5  # 5 смен по 11 ч больше суток
    sheet.append(["", "Высота потолка в попугаях", 3])
    sheet.append(["", "Площадь зоны работы роботов", 2, "га"])
    people = book["Штат"]
    people.append(["Космонавт", 3, 3, 100000, "нет"])
    out = BytesIO()
    book.save(out)

    result = upload(out.getvalue(), "мой склад.xlsx").json()
    assert result["overrides"][AREA] == 12000
    assert SHIFTS not in result["overrides"]
    messages = {(row["field"], row["status"]): row["message"] for row in result["rows"]}
    assert "в сутках 24" in messages[("Смен в сутки", "error")]
    assert "нет в расчете" in messages[("Высота потолка в попугаях", "error")]
    assert "«м2»" in messages[("Площадь зоны работы роботов", "error")]
    assert "справочнике" in messages[("Космонавт", "error")]
    assert result["staff"] and all(line["role"] != "Космонавт" for line in result["staff"])
    assert result["errors"] >= 4


def test_csv_in_cp1251_with_comma_decimals():
    text = "Параметр;Значение\nПлощадь зоны работы роботов;15 000\nПродолжительность смены;10,5\n"
    result = upload(text.encode("cp1251"), "данные.csv").json()
    assert result["overrides"][AREA] == 15000
    assert result["overrides"]["facilities.warehouse.schedule.shift_hours"] == 10.5
    assert result["staff"] is None


def test_wrong_file_is_refused_with_reason():
    response = upload(b"hello", "photo.png")
    assert response.status_code == 415
    assert "Excel" in response.json()["detail"]


def test_old_ruble_spelling_in_unit_is_accepted():
    # шаблоны, скачанные до замены "руб/кВт*ч" на "₽/кВт·ч", грузятся без ошибки
    text = "Код;Значение;Единица\neconomics.energy_price_rub_kwh;8;руб/кВт*ч\n"
    result = upload(text.encode("utf-8"), "данные.csv").json()
    assert result["overrides"]["economics.energy_price_rub_kwh"] == 8
