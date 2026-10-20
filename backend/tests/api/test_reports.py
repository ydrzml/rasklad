from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.reports import names

client = TestClient(app)


def test_pdf_report_is_a_pdf_with_a_file_name():
    response = client.post("/api/reports/pdf", json={})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    assert "attachment" in response.headers["content-disposition"]
    assert ".pdf" in response.headers["content-disposition"]


def _years_total(book, scenario: str) -> float:
    """Накопленные затраты сценария по листу «По годам»: сумма колонок людей, роботов и вложений"""
    rows = [r for r in book["По годам"].iter_rows(min_row=6, values_only=True) if r[0] == scenario]
    return sum(r[2] + r[3] + r[4] for r in rows)


def test_excel_holds_the_same_numbers_as_the_screen():
    preview = client.post("/api/calculations/preview", json={}).json()
    response = client.post("/api/reports/xlsx", json={})
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    # без прогона смены листа «Смена» нет, чувствительность есть всегда
    assert book.sheetnames == [
        "Итог",
        "По годам",
        "Затраты",
        "Что будет, если",
        "Парк и штат",
        "Цена операции",
        "Вопросы поставщику",
        "Параметры",
        "Источники",
    ]
    # итог считается формулами из листа по годам, а сами годы складываются в стоимость владения
    tco = next(r for r in book["Итог"].iter_rows(values_only=True) if str(r[0]).startswith("Стоимость владения"))
    assert all(str(cell).startswith("='По годам'!G") for cell in tco[1:4])
    for name, scenario in zip(["Без роботов", "Покупка", "Аренда"], preview["scenarios"], strict=True):
        assert abs(_years_total(book, name) - scenario["tco_rub"]) < 0.01
    # каждое значение расчета попало на лист источников
    assert book["Источники"].max_row - 5 == len(preview["sources"])


def _row(book, start: str) -> tuple:
    return next(r for r in book["Итог"].iter_rows(values_only=True) if str(r[0]).startswith(start))


def test_excel_payback_and_roi_match_the_screen_in_leasing():
    # в лизинг "вложения на старте" это только аванс, а окупаемость по ТЗ считается от всех вложений:
    # формула от листа давала 0,7 года при 1,5 на экране
    body = {"overrides": {"financing.method": 2}}
    preview = client.post("/api/calculations/preview", json=body).json()
    book = load_workbook(BytesIO(client.post("/api/reports/xlsx", json=body).content))
    purchase = next(s for s in preview["scenarios"] if s["id"] == "purchase")
    assert purchase["investment_year0_rub"] < purchase["capex_total_rub"]
    payback = _row(book, "Окупаемость по ТЗ")
    assert payback[2] == pytest.approx(purchase["payback_simple_years"])
    # ROI одной строкой в процентах и одной разами: проект, который только окупился, показывал 100% и 0%
    roi = _row(book, "ROI за 5 лет")
    assert roi[2] == pytest.approx(purchase["roi_net"])
    per_rub = _row(book, "Экономия на рубль вложений")
    assert per_rub[2] == pytest.approx(purchase["roi_tz"])
    labels = [str(r[0]) for r in book["Итог"].iter_rows(values_only=True)]
    assert not any(label in ("ROI по ТЗ", "Чистый ROI") for label in labels)


def test_excel_payback_says_never_when_first_year_loses():
    small = {
        "facilities.warehouse.operations.pallet_transport.volume_per_day": 150,
        "facilities.warehouse.schedule.shifts": 1,
        "facilities.warehouse.schedule.shift_hours": 10,
        # дорогой робот: ТО 10% цены в год больше, чем экономия на одном водителе
        "robots.ronavi-h1500.price_rub": 20_000_000,
    }
    staff = [{"id": "f", "role": "forklift_operator", "headcount": 1, "filled": 1, "salary_month": 120000}]
    body = {"overrides": small, "staff": staff}
    preview = client.post("/api/calculations/preview", json=body).json()
    purchase = next(s for s in preview["scenarios"] if s["id"] == "purchase")
    assert purchase["annual_effect_year1_rub"] < 0
    book = load_workbook(BytesIO(client.post("/api/reports/xlsx", json=body).content))
    assert _row(book, "Окупаемость по ТЗ")[2] == "не окупается"


def test_report_keeps_user_edits():
    salary = "facilities.warehouse.staff.forklift-operators.salary_month"
    base = client.post("/api/calculations/preview", json={}).json()
    edited = client.post("/api/calculations/preview", json={"overrides": {salary: 150000}}).json()
    book = load_workbook(BytesIO(client.post("/api/reports/xlsx", json={"overrides": {salary: 150000}}).content))
    assert abs(_years_total(book, "Без роботов") - edited["scenarios"][0]["tco_rub"]) < 0.01
    assert edited["scenarios"][0]["tco_rub"] != base["scenarios"][0]["tco_rub"]


def test_unknown_robot_is_not_found():
    response = client.post("/api/reports/pdf", json={"robot_id": "nope"})
    assert response.status_code in (404, 422)


def test_cumulative_costs_end_at_tco():
    for scenario in client.post("/api/calculations/preview", json={}).json()["scenarios"]:
        points = scenario["cumulative_cost_rub"]
        assert points[0] == scenario["investment_year0_rub"]
        assert abs(points[-1] - scenario["tco_rub"]) < 1
        assert abs(scenario["investment_year0_rub"] + scenario["running_cost_rub"] - scenario["tco_rub"]) < 1


def test_sources_hold_only_the_chosen_task_and_staff_names_are_russian():
    # уборка и отбор в расчете перевозки паллет не участвуют: ни на экране, ни в отчете их нет
    preview = client.post("/api/calculations/preview", json={}).json()
    tasks = {s["path"].split(".")[3] for s in preview["sources"] if ".operations." in s["path"]}
    assert tasks == {"pallet_transport"}
    assert all(".operations.pallet_transport." in p or ".operations." not in p for p in preview["weak_value_paths"])

    path = "facilities.warehouse.staff.pickers.salary_month"
    assert names.name(path, {}, None, {"pickers": "Отборщик"}) == "Оклад в месяц, отборщик"


def test_report_time_is_moscow_and_says_so():
    book = load_workbook(BytesIO(client.post("/api/reports/xlsx", json={}).content))
    cells = [str(cell) for row in book["Итог"].iter_rows(values_only=True) for cell in row if cell]
    assert any("расчет от" in cell and "мск" in cell for cell in cells)


def test_report_refuses_robot_that_cannot_lift_the_load_the_user_entered():
    """Груз 3000 кг Ronavi H1500 не поднимает: подбор его исключает, и отчет по нему не собирается."""
    import pytest

    from app.reports import content
    from app.schemas.economics import CalculationRequest

    load = "facilities.warehouse.operations.pallet_transport.load_kg"
    with pytest.raises(content.NotFeasible, match="не подходит для задачи"):
        content.build(CalculationRequest(overrides={load: 3000}))


def test_case_report_has_its_own_name_and_no_risk_about_undrawn_plan():
    """Отчет готового разбора: имя по разбору, а не одно на все три, и без пункта "План вы не чертили":
    план в разборе типовой у всех, и страница разбора этот пункт не пишет"""
    from app.reports import content
    from app.services import demo_cases

    request = demo_cases.get("cleaning").request
    plain = content.build(request)
    case = content.build(request, case_id="cleaning")
    assert any(risk.startswith(content.NOT_DRAWN) for risk in plain.risks)
    assert not any(risk.startswith(content.NOT_DRAWN) for risk in case.risks)
    assert case.file_stem.startswith("razbor-cleaning-")
    response = client.post("/api/reports/xlsx?case=nope", json={})
    assert response.status_code == 404
