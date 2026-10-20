"""Сводка шага экономики, чувствительность и журнал смены. Числа сверены с расчетом руками
на демо-складе без плана (маршрут 75 м по площади) при режиме 2 смены по 11 ч, 365 дней: формула,
спрос с резервом 1900 / 22 * 1,5 * 1,15 = 148,98 в час, 148,98 / 15,88 = 9,38, вверх 10 роботов."""

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

client = TestClient(app)


def post(url: str, **body):
    response = client.post(url, json=body)
    assert response.status_code == 200, response.text
    return response


def test_tco_parts_add_up_to_tco():
    result = post("/api/calculations/preview").json()
    for scenario in result["scenarios"]:
        assert sum(part["rub"] for part in scenario["tco_parts"]) == pytest.approx(scenario["tco_rub"])
    purchase = result["scenarios"][1]
    parts = {part["id"]: part["rub"] for part in purchase["tco_parts"]}
    # вложения на старте: роботы 27.0 + зарядки 3 * 0.45 (10 роботов по 4 на станцию) + ПО 1.5
    # + интеграция 5.0 + инфраструктура 3.0 + пусконаладка 2.7 = 40.55 с НДС, без НДС 22% / 1.22 = 33.2377,
    # + переобучение на присмотр 0.0127 (курсы без НДС) = 33.2504, резерв 10% сверху: 36.575 млн
    assert parts["start"] == pytest.approx(36.575e6, abs=0.01e6)
    # присмотр: 10 * 1.5 / 60 = 0.25 поста * (22 * 365 / 1972 * 1.25 = 5.09) = 1.2725 ставки
    # * 1 880 640 рублей * (1 + 1.07 + ... + 1.07^4)
    assert parts["operators"] == pytest.approx(1.2725 * 1_880_640 * 5.750739, rel=1e-3)


def test_amortization_is_linear_on_service_life():
    purchase, raas = post("/api/calculations/preview").json()["scenarios"][1:]
    # покупка: 36.575 млн / 5 лет службы = 7.32 млн в год
    assert purchase["amortization"]["method"] == "линейная"
    assert purchase["amortization"]["years"] == 5
    assert purchase["amortization"]["per_year_rub"] == pytest.approx(36.575e6 / 5, abs=0.01e6)
    # аренда базово без выкупа: ((5 + 3 + 2.7) / 1.22 + 0.0127) * 1.1 = 9.662 млн / 3 года договора
    assert raas["amortization"]["years"] == 3
    assert raas["amortization"]["per_year_rub"] == pytest.approx(9.662e6 / 3, abs=0.01e6)
    assert raas["amortization"]["after_buyout_base_rub"] is None
    # с выкупом: 10 * 2.7 * 0.4 / 1.22 = 8.852 млн / 2 оставшихся года службы
    raas = post("/api/calculations/preview", raas_buyout=True).json()["scenarios"][2]
    assert raas["amortization"]["after_buyout_base_rub"] == pytest.approx(10.8e6 / 1.22)
    assert raas["amortization"]["after_buyout_per_year_rub"] == pytest.approx(5.4e6 / 1.22)
    assert purchase["effect_total_rub"] / purchase["investment_total_rub"] == pytest.approx(purchase["roi_tz"])


def test_sources_have_units():
    sources = {s["path"]: s for s in post("/api/calculations/preview").json()["sources"]}
    assert sources["robots.ronavi-h1500.price_rub"]["unit"] == "₽"
    assert sources["robots.ronavi-h1500.raas.fee_rub_month"]["unit"] == "₽ в месяц за робота"
    assert sources["facilities.warehouse.staff.forklift-operators.salary_month"]["unit"] == "₽ в месяц"


def _cell(result, param: str, delta: float, scenario: str = "purchase"):
    row = next(p for p in result["params"] if p["id"] == param)
    cell = next(c for c in row["cells"] if c["delta"] == delta)
    return cell, next(o for o in cell["outcomes"] if o["scenario_id"] == scenario)


def test_sensitivity_matches_hand_count():
    result = post("/api/calculations/sensitivity").json()
    assert result["deltas"] == [-0.2, -0.1, 0.0, 0.1, 0.2]
    assert {p["id"] for p in result["params"]} >= {"wages", "robot_price", "volume"}
    base = post("/api/calculations/preview").json()["scenarios"][1]
    cell, same = _cell(result, "robot_price", 0.0)
    assert same["tco_rub"] == pytest.approx(base["tco_rub"])
    # Цена робота +10%, все без НДС 22%: 10 роботов +2.7 / 1.22 млн, пусконаладка +0.27 / 1.22, резерв 10%
    # от них: (2.97 / 1.22) * 1.1 = 2.677869; ТО 10% цены в год (у производителя, #65):
    # +0.27 / 1.22 млн * (1 + 1.04 + ... + 1.04^4 = 5.416323) = 1.198694. Итого +3.876563 млн
    _, up = _cell(result, "robot_price", 0.1)
    assert up["tco_rub"] - same["tco_rub"] == pytest.approx(3_876_563.2, abs=1)

    # Объем -10%: спрос 148.98 * 0.9 = 134.08 в час, 134.08 / 15.88 = 8.44, вверх 9 роботов
    cell, _ = _cell(result, "volume", -0.1)
    assert cell["fleet"] == 9
    # плата за аренду покупку не меняет
    _, fee = _cell(result, "raas_fee", 0.2)
    assert fee["tco_rub"] == pytest.approx(same["tco_rub"])
    # дороже люди, быстрее окупается
    _, cheap = _cell(result, "wages", -0.2)
    _, dear = _cell(result, "wages", 0.2)
    assert cheap["payback_years"] > same["payback_years"] > dear["payback_years"]


def test_shift_log_is_csv_with_every_segment():
    response = post("/api/simulation/log.csv", fleet=9)
    assert response.headers["content-type"].startswith("text/csv")
    assert ".csv" in response.headers["content-disposition"]
    lines = response.content.decode("utf-8-sig").splitlines()
    head = lines[0].split(";")
    assert head[0] == "робот" and "действие" in head and "x начала, м" in head
    rows = [line.split(";") for line in lines[1:]]
    assert len(rows) > 100
    assert {row[0] for row in rows} == {str(n) for n in range(1, 10)}
    assert {row[5] for row in rows} <= {"едет", "грузит", "ждет", "заряжается", "без задания"}


def test_report_holds_shift_and_sensitivity():
    book = load_workbook(BytesIO(post("/api/reports/xlsx", use_simulation=True).content))
    assert "Что будет, если" in book.sheetnames
    assert "Смена" in book.sheetnames
    shift = {row[0]: row[1] for row in book["Смена"].iter_rows(min_row=6, values_only=True)}
    assert shift["Роботов в прогоне"] == 9

    assert shift["Единица задачи"] == "паллет"
    assert shift["Сделано за смену"] > 0
    units = [row[3] for row in book["Источники"].iter_rows(min_row=6, values_only=True)]
    assert "₽" in units
    pdf = post("/api/reports/pdf", use_simulation=True).content
    assert pdf.startswith(b"%PDF")
