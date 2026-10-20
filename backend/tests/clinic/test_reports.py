"""PDF и Excel медучреждения через сервер и сохранение итога в проект, как у склада."""

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

client = TestClient(app)
REQUEST = {"operation_ids": ["floor_delivery", "cleaning"]}
WHO = {"email": "clinic@example.com", "password": "correct-horse"}


def economics(api, body):
    response = api.post("/api/clinic/economics", json=body)
    assert response.status_code == 200, response.text
    return response.json()


def test_clinic_pdf_is_built_by_the_server():
    response = client.post("/api/reports/clinic/pdf", json=REQUEST)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    assert "otsenka-clinic" in response.headers["content-disposition"]


def test_clinic_excel_has_the_same_numbers_as_the_screen():
    answer = economics(client, REQUEST)
    response = client.post("/api/reports/clinic/xlsx", json=REQUEST)
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames[:2] == ["Итог", "По годам"]
    assert {"Парк и смена", "Люди", "Ограничения", "Источники"} <= set(book.sheetnames)
    baseline = next(one for one in answer["scenarios"] if one["id"] == "baseline")
    assert book["По годам"].cell(row=7, column=2).value == pytest.approx(baseline["years_rub"][0], rel=1e-12)
    assert book["Источники"].max_row - 5 == len(answer["sources"])


def test_clinic_report_names_the_reason_when_the_plan_cannot_be_counted():
    building = client.get("/api/clinic/template").json()
    body = {"operation_ids": ["cleaning"], "building": building, "cleaner_id": "446c5207-a099-45e0-b615-afd60de08589"}
    response = client.post("/api/reports/clinic/pdf", json=body)
    assert response.status_code == 422
    assert "мм" in response.text


def test_clinic_result_saved_to_a_project_gives_a_report_and_recalculates(db_client):
    if db_client.post("/api/auth/login", json=WHO).status_code != 200:
        assert db_client.post("/api/auth/register", json=WHO).status_code in (200, 201)
    answer = economics(db_client, REQUEST)
    state = {"facilityId": "clinic", "taskIds": REQUEST["operation_ids"], "clinicRequest": REQUEST}
    created = db_client.post(
        "/api/projects",
        json={"name": "Медучреждение", "facility_type": "clinic", "state": state, "result": answer, "note": ""},
    )
    assert created.status_code == 201, created.text
    project = created.json()
    number = project["current"]["number"]
    for kind, start in (("pdf", b"%PDF"), ("xlsx", b"PK")):
        report = db_client.get(f"/api/projects/{project['id']}/versions/{number}/report/{kind}")
        assert report.status_code == 200, report.text
        assert report.content.startswith(start)
    again = db_client.post(f"/api/projects/{project['id']}/versions/{number}/recalculate")
    assert again.status_code == 201, again.text
    assert again.json()["result"]["horizon_years"] == answer["horizon_years"]
    listed = next(one for one in db_client.get("/api/projects").json() if one["id"] == project["id"])
    assert listed["figures"]["tco_rub"] is not None
