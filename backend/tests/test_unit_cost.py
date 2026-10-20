"""Стоимость одной операции: все расходы сценария за горизонт на число операций за горизонт."""

import pytest

from app.schemas.economics import CalculationRequest
from app.services.calculation import calculate
from app.services.unit_cost import compute


@pytest.fixture(scope="module")
def result():
    return calculate(CalculationRequest())


def test_every_scenario_gets_a_price_per_operation(result):
    assert len(result.unit_costs) == 1
    cost = result.unit_costs[0]
    assert cost.operation_id == "pallet_transport"
    assert cost.per == "перевозку паллеты"
    assert {one.id for one in cost.scenarios} == {"baseline", "purchase", "raas"}
    assert all(one.rub_per_unit > 0 for one in cost.scenarios)


def test_price_is_scenario_total_divided_by_operations_in_horizon(result):
    """Одна задача: общее на объект целиком уходит в нее, значит цена это стоимость владения на операции"""
    cost = result.unit_costs[0]
    for scenario in result.scenarios:
        one = next(row for row in cost.scenarios if row.id == scenario.id)
        assert one.rub_per_unit * cost.operations_horizon == pytest.approx(scenario.tco_rub, rel=1e-6)


def test_volume_is_per_day_times_working_days_times_years(result):
    # 2000 паллет в сутки, 365 рабочих дней в году у склада по умолчанию, горизонт расчета
    assert result.unit_costs[0].operations_horizon == pytest.approx(2000 * 365 * result.horizon_years)


def test_shared_costs_are_split_in_proportion_to_task_cost(result):
    """Свои деньги задач плюс доля общего дают весь сценарий, ничего не теряется и не считается дважды"""
    facility = {"operations": [{"id": "pallet_transport", "volume_per_day": 100}], "schedule": {"days_per_year": 365}}
    tasks = result.tasks
    first = compute(tasks, result.shared, facility, result.horizon_years)[0]
    purchase = next(row for row in first.scenarios if row.id == "purchase")
    task_tco = sum(part.tco_rub for task in tasks for part in task.scenarios if part.id == "purchase")
    shared_tco = next(part.tco_rub for part in result.shared.scenarios if part.id == "purchase")
    assert purchase.rub_per_unit * first.operations_horizon == pytest.approx(task_tco + shared_tco, rel=1e-6)


def test_task_without_volume_is_skipped(result):
    facility = {"operations": [{"id": "pallet_transport", "volume_per_day": 0}], "schedule": {"days_per_year": 365}}
    assert compute(result.tasks, result.shared, facility, result.horizon_years) == []


def test_nothing_to_split_gives_no_rows():
    assert compute([], None, {"operations": []}, 5) == []


def test_supplier_questions_come_with_the_calculation(result):
    topics = [one.topic for one in result.supplier_questions]
    assert 8 <= len(topics) <= 10
    assert {"Зарядка", "Связь с WMS", "Сервис", "Сроки", "Обучение", "Гарантия"} <= set(topics)
    assert all(one.question.endswith(("?", ".")) for one in result.supplier_questions)


def test_charging_question_names_the_fleet_and_chargers_of_this_calculation(result):
    charging = next(one for one in result.supplier_questions if one.topic == "Зарядка")
    assert f"парке в {result.sizing.fleet} шт." in charging.question
    assert f"В расчете их {result.sizing.chargers}." in charging.question


def test_excel_report_has_unit_cost_and_supplier_sheets():
    from io import BytesIO

    from openpyxl import load_workbook

    from app.reports import content, excel

    book = load_workbook(BytesIO(excel.render(content.build(CalculationRequest()))))
    assert "Цена операции" in book.sheetnames
    assert "Вопросы поставщику" in book.sheetnames
    rows = list(book["Цена операции"].iter_rows(min_row=5, values_only=True))
    assert rows[1][0] == "₽ за перевозку паллеты"
    assert all(isinstance(value, float) for value in rows[1][2:5])
    assert book["Вопросы поставщику"].max_row >= 5 + 8


def test_pdf_report_has_unit_cost_and_supplier_sections():
    from app.reports import content, pdf

    body = pdf.render(content.build(CalculationRequest()))
    assert body.startswith(b"%PDF")
    assert len(body) > 10_000
