"""Нормативы в админке: администратор видит и правит взносы, рабочие часы, резервы, цены и
значения по умолчанию параметров объекта (ТЗ, п. 3.1.4).

Файл config/model.yaml остается базой по умолчанию со своими источниками. Правка администратора
лежит в базе (norm_overrides) и накладывается на модель при каждом расчете, перезапуск не нужен.
Какие значения открыты, в каком разделе, с какими границами и что они меняют, говорит config/norms.yaml.

Как расчет узнает о правке. Перед чтением модели спрашиваем у базы номер последней записи журнала
нормативов, это один маленький запрос. Поменялся, значит правки перечитываем и сбрасываем кеш
поиска парка. Так правку видят и процессы, которые гоняют прогон смены параллельно: у них своя память.
"""

from __future__ import annotations

import csv
import hashlib
import io
import math
import re
from dataclasses import dataclass
from datetime import datetime
from functools import lru_cache

import yaml
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.engine.model_config import TRUST_LEVELS, Provenance, overlay, value_by_path
from app.reports import names
from app.reports.cells import defuse
from app.schemas import norms as schemas
from app.schemas.data_import import DataImportRow
from app.services import data_import
from app.services.catalog_export import cell, uncell
from app.settings import settings
from app.storage.models import NormChange, NormOverride, User

HEAD = ["Код", "Название", "Значение", "Единица", "Источник", "Дата", "Оценка"]
DATE = re.compile(r"\d{4}-\d{2}(-\d{2})?")
FROM_FILE = "Значение из файла модели"


class NotFound(KeyError):
    """Такого норматива нет или он не открыт для правки."""


class Invalid(ValueError):
    """Значение не подходит: не число, за границами, невозможный режим работы."""


@dataclass(frozen=True)
class Rule:
    pattern: str
    group: str
    group_name: str
    name: str | None
    min: float
    max: float
    whole: bool
    effect: str


@dataclass(frozen=True)
class Saved:
    value: float
    source: str
    date: str | None
    trust: str
    by: str
    at: datetime


@lru_cache(maxsize=1)
def rules() -> tuple[list[schemas.NormGroup], list[Rule]]:
    raw = yaml.safe_load((settings.config_dir / "norms.yaml").read_text(encoding="utf-8"))
    groups, found = [], []
    for group in raw["groups"]:
        groups.append(schemas.NormGroup(id=group["id"], name=group["name"]))
        for norm in group["norms"]:
            found.append(
                Rule(
                    pattern=norm["path"],
                    group=group["id"],
                    group_name=group["name"],
                    name=norm.get("name"),
                    min=float(norm["min"]),
                    max=float(norm["max"]),
                    whole=bool(norm.get("whole", False)),
                    effect=norm["effect"],
                )
            )
    return groups, found


@lru_cache(maxsize=1)
def _typical() -> dict[str, tuple[float, float]]:
    """Диапазоны датасета из полей формы параметров: за ними предупреждаем, но значение принимаем."""
    raw = yaml.safe_load((settings.config_dir / "parameters.yaml").read_text(encoding="utf-8"))
    return {
        field["path"]: (float(field["min"]), float(field["max"]))
        for fields in raw.values()
        for field in fields
        if "min" in field and "max" in field
    }


def _matches(pattern: str, path: str) -> bool:
    wanted, parts = pattern.split("."), path.split(".")
    return len(wanted) == len(parts) and all(w in ("*", p) for w, p in zip(wanted, parts, strict=True))


def _base() -> tuple[dict, dict[str, Provenance]]:
    from app.services import calculation  # расчет сам зовет нормативы, поэтому модель берем при вызове

    return calculation._model()


def open_paths() -> dict[str, Rule]:
    """Открытые нормативы по порядку разделов из config/norms.yaml. Путь без числа в модели не берем."""
    model, provenance = _base()
    found: dict[str, Rule] = {}
    for rule in rules()[1]:
        for path in provenance:
            if _matches(rule.pattern, path) and isinstance(value_by_path(model, path), (int, float)):
                found.setdefault(path, rule)
    return found


def _name(model: dict, path: str, rule: Rule) -> str:
    if rule.name:
        return rule.name
    parts = path.split(".")
    facility = next((f for f in model["facilities"] if f["id"] == parts[1]), None) if parts[0] == "facilities" else None
    operations = {op["id"]: op["name"] for op in facility.get("operations", [])} if facility else {}
    roles = model.get("roles", {})
    staff = (
        {line["id"]: roles.get(line.get("role"), {}).get("name", line["id"]) for line in facility.get("staff", [])}
        if facility
        else {}
    )
    name = names.name(path, {}, operations, staff)
    if facility and len(model["facilities"]) > 1:
        name += f", {facility.get('name', facility['id']).lower()}"
    return name


# Что правил администратор. Отметка: номер последней записи журнала, по ней понимаем, что пора перечитать
_state: tuple[int | None, dict[str, Saved]] | None = None
_applied: tuple[dict, int | None, tuple[dict, dict[str, Provenance]]] | None = None


def forget() -> None:
    global _state, _applied
    _state, _applied = None, None


def _read_saved(session: Session) -> dict[str, Saved]:
    return {
        row.path: Saved(row.value, row.source, row.source_date, row.trust, row.user_email, row.updated_at)
        for row in session.scalars(select(NormOverride))
    }


def _current() -> tuple[int | None, dict[str, Saved]]:
    global _state
    from app.services import selection  # сессии каталога одни на подбор, цены и нормативы, их подменяют тесты

    with selection.sessions() as session:
        stamp = session.scalar(select(func.max(NormChange.id)))
        if _state is not None and _state[0] == stamp:
            return _state
        saved = _read_saved(session)
    if _state is not None:
        # Поиск парка держит ответы в памяти по правкам пользователя, а нормативы в ключ не входят
        from app.services import simulation

        simulation.fleet_search.cache_clear()
    _state = (stamp, saved)
    return _state


def provenance_of(saved: Saved) -> Provenance:
    when = saved.at.strftime("%d.%m.%Y") if saved.at else ""
    who = f" {saved.by}" if saved.by else ""
    return Provenance(f"Норматив, правил администратор{who} {when}: {saved.source}".strip(), saved.date, saved.trust)


def apply(model: dict, provenance: dict[str, Provenance]) -> tuple[dict, dict[str, Provenance]]:
    """Модель с правками администратора. Пока правок нет, отдаем ту же модель без копии."""
    global _applied
    stamp, saved = _current()
    if _applied is not None and _applied[0] is model and _applied[1] == stamp:
        return _applied[2]
    result = overlay(model, provenance, {path: (s.value, provenance_of(s)) for path, s in saved.items()})
    _applied = (model, stamp, result)
    return result


def fingerprint() -> str:
    """Отпечаток правок для версии данных. Пусто, пока правок нет: старые проекты остаются на тех же данных."""
    saved = _current()[1]
    if not saved:
        return ""
    text = "\n".join(f"{path}={s.value!r}" for path, s in sorted(saved.items()))
    return hashlib.sha256(text.encode()).hexdigest()[:12]


def listing(session: Session) -> schemas.NormList:
    model, provenance = _base()
    saved = _read_saved(session)
    typical = _typical()
    items = []
    for path, rule in open_paths().items():
        file_value, file_source = float(value_by_path(model, path)), provenance[path]
        mine = saved.get(path)
        now = provenance_of(mine) if mine else file_source
        low, high = typical.get(path, (None, None))
        items.append(
            schemas.Norm(
                code=path,
                name=_name(model, path, rule),
                group=rule.group,
                group_name=rule.group_name,
                unit=names.unit(path),
                value=mine.value if mine else file_value,
                source=mine.source if mine else file_source.source,
                date=now.date,
                trust=now.trust,
                effect=rule.effect,
                min=rule.min,
                max=rule.max,
                typical_min=low,
                typical_max=high,
                whole=rule.whole,
                file=schemas.NormSource(
                    value=file_value, source=file_source.source, date=file_source.date, trust=file_source.trust
                ),
                edited=schemas.NormEdited(by=mine.by, at=mine.at) if mine else None,
            )
        )
    return schemas.NormList(groups=rules()[0], items=items)


def changes(session: Session, limit: int = 100, path: str | None = None) -> list[schemas.NormChange]:
    query = select(NormChange).order_by(NormChange.at.desc(), NormChange.id.desc()).limit(limit)
    if path:
        query = query.where(NormChange.path == path)
    return [schemas.NormChange.model_validate(row, from_attributes=True) for row in session.scalars(query)]


def _rule(path: str) -> Rule:
    rule = open_paths().get(path)
    if rule is None:
        raise NotFound(path)
    return rule


def check(path: str, value: float, others: dict[str, float] | None = None) -> None:
    """Правила значения: число, целое где нужно, в жестких границах, и объект с ним возможен."""
    from app.services import calculation

    rule = _rule(path)
    if not math.isfinite(value):
        raise Invalid("нужно обычное число")
    if rule.whole and value != int(value):
        raise Invalid("нужно целое число")
    if not rule.min <= value <= rule.max:
        raise Invalid(f"допустимо от {_show(rule.min)} до {_show(rule.max)}")
    model, provenance = apply(*_base())
    trial = {**(others or {}), path: value}
    changed, _ = overlay(model, provenance, {p: (v, provenance[p]) for p, v in trial.items()})
    try:
        calculation.check_possible(changed)
    except calculation.Impossible as error:
        # Текст расчета отправляет на шаг параметров, а тут правят в админке
        raise Invalid(str(error).replace(" на шаге параметров", "")) from error


def _check_source(edit: schemas.NormEdit) -> None:
    if not edit.source.strip():
        raise Invalid("нужен источник или объяснение, откуда значение")
    if edit.date and not DATE.fullmatch(edit.date):
        raise Invalid("дата пишется как ГГГГ-ММ или ГГГГ-ММ-ДД")
    if edit.trust != "F" and not edit.date:
        raise Invalid(f"у значения с оценкой {edit.trust} должна быть дата источника")


def save(session: Session, user: User, path: str, edit: schemas.NormEdit, action: str = "edit", note: str = "") -> bool:
    """Записать правку и журнал. Значение и источник как в файле снимают правку. False, если ничего не поменялось."""
    _rule(path)
    _check_source(edit)
    check(path, edit.value)
    return _write(session, user, path, edit, action, note)


def _write(session: Session, user: User | None, path: str, edit: schemas.NormEdit, action: str, note: str) -> bool:
    model, provenance = _base()
    file = schemas.NormEdit(
        value=float(value_by_path(model, path)),
        source=provenance[path].source,
        date=provenance[path].date,
        trust=provenance[path].trust,
    )
    row = session.get(NormOverride, path)
    before = (
        schemas.NormEdit(value=row.value, source=row.source, date=row.source_date, trust=row.trust) if row else file
    )
    after = edit.model_copy(update={"source": edit.source.strip()})
    if after == before:
        return False
    diff = {
        key: [getattr(before, key), getattr(after, key)]
        for key in ("value", "source", "date", "trust")
        if getattr(before, key) != getattr(after, key)
    }
    if after == file:
        session.delete(row)
        action, note = ("reset" if action == "edit" else action), note or FROM_FILE
    else:
        if row is None:
            row = NormOverride(path=path)
            session.add(row)
        row.value, row.source, row.source_date, row.trust = after.value, after.source, after.date, after.trust
        row.user_id, row.user_email = (user.id, user.email) if user else (None, "")
        row.updated_at = func.now()
    _journal(session, user, action, path, diff, note)
    return True


def reset(session: Session, user: User, path: str) -> bool:
    """Вернуть значение из файла модели. False, если правки и не было."""
    _rule(path)
    row = session.get(NormOverride, path)
    if row is None:
        return False
    model, provenance = _base()
    source = provenance[path]
    diff = {
        "value": [row.value, float(value_by_path(model, path))],
        "source": [row.source, source.source],
        "date": [row.source_date, source.date],
        "trust": [row.trust, source.trust],
    }
    session.delete(row)
    _journal(session, user, "reset", path, {k: v for k, v in diff.items() if v[0] != v[1]}, FROM_FILE)
    return True


def _journal(session: Session, user: User | None, action: str, path: str, diff: dict, note: str) -> None:
    model, _ = _base()
    session.add(
        NormChange(
            user_id=user.id if user else None,
            user_email=user.email if user else "",
            action=action,
            path=path,
            name=_name(model, path, _rule(path)),
            changes=diff,
            note=note,
        )
    )


# Выгрузка и загрузка файлом


def _rows(session: Session) -> list[list[str]]:
    return [
        [
            norm.code,
            norm.name,
            _show(norm.value),
            norm.unit,
            norm.source,
            norm.date or "",
            norm.trust,
        ]
        for norm in listing(session).items
    ]


def export_csv(session: Session) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";", lineterminator="\r\n")
    writer.writerow(HEAD)
    writer.writerows([[cell(value) for value in row] for row in _rows(session)])
    # С BOM Excel сам поймет, что файл в UTF-8
    return ("﻿" + out.getvalue()).encode("utf-8")


def export_xlsx(session: Session) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Нормативы"
    sheet.append(HEAD)
    for row in _rows(session):
        # Значение числом, чтобы в Excel его можно было считать, остальное текстом: дату Excel не переделает
        sheet.append([row[0], row[1], _number_or_text(row[2]), *row[3:]])
        for index in (5,):
            sheet.cell(sheet.max_row, index + 1).number_format = "@"
    for column, width in zip("ABCDEFG", (48, 42, 14, 16, 80, 12, 8), strict=True):
        sheet.column_dimensions[column].width = width
    defuse(book)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _number_or_text(text: str) -> float | str:
    number = data_import._number(text)
    return text if number is None else number


def import_file(session: Session, user: User, content: bytes, filename: str, dry_run: bool) -> schemas.NormImportResult:
    """Разобрать файл в формате выгрузки и, если не dry_run, записать правки. Отчет по каждой строке.

    Строка с ошибкой не мешает остальным. Значение и источник как в файле модели снимают правку,
    как в текущем значении идут строкой "без изменений": выгрузка, загруженная обратно, ничего не меняет.
    """
    if len(content) > data_import.MAX_BYTES:
        raise data_import.TooLarge
    tables = data_import._read(content, filename)
    report: list[DataImportRow] = []
    accepted: dict[str, float] = {}
    edits: list[tuple[str, schemas.NormEdit]] = []
    current = {norm.code: norm for norm in listing(session).items}
    found_head = False
    for sheet, table in tables:
        head_at = next(
            (i for i, row in enumerate(table) if _col(row, "код") is not None and _col(row, "значен") is not None),
            None,
        )
        if head_at is None:
            continue
        found_head = True
        head = table[head_at]
        columns = {key: _col(head, word) for key, word in COLUMNS.items()}
        for index, cells in enumerate(table[head_at + 1 :], start=head_at + 2):
            if not any(str(c).strip() for c in cells):
                continue
            row = {key: uncell(data_import._cell(cells, at)) or "" for key, at in columns.items()}
            outcome = _import_row(row, current, accepted)
            status, message, edit = outcome
            report.append(
                DataImportRow(
                    sheet=sheet,
                    row=index,
                    field=row["code"] or "(нет кода)",
                    value=row["value"],
                    status=status,
                    message=message,
                )
            )
            if edit is not None:
                edits.append((row["code"], edit))
    if not found_head:
        raise data_import.WrongFile('Не нашли строку заголовков: нужны колонки "Код" и "Значение", как в выгрузке')
    applied = 0
    for path, edit in edits:
        if dry_run:
            applied += 1
        elif _write(session, user, path, edit, "import", f"файл {filename}"[:200]):
            applied += 1
    if not dry_run:
        session.commit()
    return schemas.NormImportResult(
        dry_run=dry_run,
        rows=report,
        applied=applied,
        unchanged=sum(1 for r in report if r.status == "ok" and r.message.startswith("Без изменений")),
        warnings=sum(1 for r in report if r.status == "warn"),
        errors=sum(1 for r in report if r.status == "error"),
    )


COLUMNS = {
    "code": "код",
    "value": "значен",
    "unit": "единиц",
    "source": "источник",
    "date": "дата",
    "trust": "оценк",
}


def _col(head: list[str], word: str) -> int | None:
    return next((i for i, text in enumerate(head) if word in data_import._norm(text)), None)


def _import_row(
    row: dict[str, str], current: dict[str, schemas.Norm], accepted: dict[str, float]
) -> tuple[str, str, schemas.NormEdit | None]:
    code = row["code"].strip()
    norm = current.get(code)
    if norm is None:
        return "error", "Такого норматива нет или он не открыт для правки: код как в выгрузке", None
    notes: list[str] = []
    value = data_import._number(row["value"])
    if value is None:
        return "error", "Значение должно быть числом", None
    unit, ours = data_import._norm(row["unit"]), data_import._norm(norm.unit)
    if unit and unit != ours:
        if unit == "%" and ours.startswith("доля"):
            value = round(value / 100, 12)
            notes.append(f"перевели из процентов: {_show(value)}")
        else:
            return "error", f'Единица "{row["unit"]}", а у норматива "{norm.unit}". Пересчитайте в нашу единицу', None
    elif not unit:
        notes.append(f"единица не указана, считаем в {norm.unit or 'штуках'}")
    trust = row["trust"].strip().upper() or norm.trust
    if trust not in TRUST_LEVELS or len(trust) != 1:
        return "error", "Оценка пишется одной буквой от S до F", None
    date = _date(row["date"])
    if row["date"].strip() and date is None:
        return "error", "Дата пишется как ГГГГ-ММ или ГГГГ-ММ-ДД", None
    source = row["source"].strip() or norm.source
    edit = schemas.NormEdit(value=value, source=source, date=date, trust=trust)
    try:
        _check_source(edit)
        check(code, value, accepted)
    except Invalid as error:
        return "error", sentence(str(error)), None
    if norm.typical_min is not None and not norm.typical_min <= value <= (norm.typical_max or math.inf):
        notes.append(f"за диапазоном датасета {_show(norm.typical_min)}-{_show(norm.typical_max or 0)}")
    accepted[code] = value
    same_as_now = schemas.NormEdit(value=norm.value, source=norm.source, date=norm.date, trust=norm.trust)
    if edit == same_as_now:
        return "ok", "Без изменений", None
    file = norm.file
    if edit == schemas.NormEdit(value=file.value, source=file.source, date=file.date, trust=file.trust):
        notes.insert(0, "совпадает с файлом модели, правка администратора снимется")
    status = "warn" if notes else "ok"
    return status, sentence("; ".join(notes)), edit


def _date(raw: str) -> str | None:
    text = raw.strip()
    if not text:
        return None
    if DATE.fullmatch(text):
        return text
    # Excel отдает дату датой: 2026-09-25 00:00:00
    match = re.fullmatch(r"(\d{4}-\d{2}-\d{2})[ T]00:00:00", text)
    if match:
        return match.group(1)
    match = re.fullmatch(r"(\d{2})\.(\d{2})\.(\d{4})", text)
    if match:
        return f"{match.group(3)}-{match.group(2)}-{match.group(1)}"
    return None


def sentence(text: str) -> str:
    """Первая буква заглавная, остальные как были: в сообщениях бывают ГГГГ-ММ и оценки буквами."""
    return text[:1].upper() + text[1:]


def _show(value: float) -> str:
    """Число для человека и для файла: с запятой, без экспоненты и лишних нулей. 2979000, 0,151"""
    text = f"{value:.10f}".rstrip("0").rstrip(".")
    return text.replace(".", ",")
