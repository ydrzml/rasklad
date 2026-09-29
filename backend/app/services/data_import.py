"""Данные объекта из файла: шаблон Excel и разбор того, что человек заполнил (ТЗ, импорт Excel/CSV).

Проверяем каждую строку отдельно и говорим, что с ней сделали. Ошибка в одной строке не мешает
остальным: подставим то, что прочиталось. Правила те же, что при вводе руками:
- выход за диапазон датасета это предупреждение, значение подставим и расчет пометим;
- невозможное (меньше нуля, больше 24 часов в сутки) это ошибка, такое не подставляем;
- единица в файле должна совпадать с нашей, пересчитывать чужие единицы молча мы не беремся.
"""

from __future__ import annotations

import csv
import io
import itertools
import math
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from app.reports.cells import defuse
from app.schemas.catalog import Parameter
from app.schemas.data_import import DataImportResult, DataImportRow, TemplateRequest
from app.schemas.staff import Role, StaffLine
from app.services import calculation
from app.services import catalog as catalog_service
from app.services import staff as staff_service

MAX_BYTES = 2 * 1024 * 1024
# В шаблоне десятки строк и семь колонок, больше читать незачем
MAX_ROWS = 1000
MAX_COLUMNS = 20

OBJECT = "Объект"
STAFF = "Штат"
ROLES = "Роли"
OBJECT_HEAD = ["Код", "Параметр", "Значение", "Единица", "Допустимо от", "Допустимо до", "Подсказка"]
STAFF_HEAD = ["Роль", "Мест в штате", "Занято людьми", "Оклад в месяц, руб", "Люди подрядчика (да/нет)"]
# так заканчивается замечание о поле, которое в файле пустое или его нет: значение осталось прежним
KEPT = "Оставили значение по умолчанию"
YES = {"да", "yes", "true", "1", "+", "подрядчик"}
NO = {"нет", "no", "false", "0", "-", ""}


class Semicolon(csv.excel):
    delimiter = ";"


class WrongFile(ValueError):
    """Файл не Excel и не CSV или его не прочитать."""


class TooLarge(ValueError):
    pass


# Прежние названия полей: шаблоны, скачанные до переименования, загружаются как раньше
OLD_LABELS = {"Площадь активной зоны": "Площадь зоны работы роботов"}


def _client_fields(facility_id: str, operation_ids: list[str]) -> list[Parameter]:
    # В файле только то, что клиент знает про объект: наши допущения правятся на шаге экономики
    return [field for field in catalog_service.parameters(facility_id, operation_ids) if not field.ours]


def template(request: TemplateRequest) -> bytes:
    """Шаблон, заполненный текущими значениями: поправить проще, чем вбивать с нуля."""
    fields = _client_fields(request.facility_id, request.operation_ids)
    roles = staff_service.roles()
    names = {role.id: role.name for role in roles}
    lines = request.staff
    if lines is None:
        lines = list(staff_service.form(request.facility_id, request.operation_ids).lines)

    book = Workbook()
    sheet = book.active
    sheet.title = OBJECT
    sheet.append(OBJECT_HEAD)
    for field in fields:
        value = request.overrides.get(field.path, field.value)
        sheet.append([field.path, field.label, value, field.unit, field.min, field.max, field.hint])
    _style(sheet, [28, 38, 14, 12, 14, 14, 60])

    people = book.create_sheet(STAFF)
    people.append(STAFF_HEAD)
    for line in lines:
        people.append(
            [
                names.get(line.role, line.role),
                line.headcount,
                line.filled,
                line.salary_month,
                "да" if line.contractor else "нет",
            ]
        )
    _style(people, [36, 14, 16, 20, 24])

    listed = book.create_sheet(ROLES)
    listed.append(["Роль", "Что делает"])
    for role in roles:
        listed.append([role.name, role.what])
    _style(listed, [36, 80])
    # роль выбирают из списка: так ее точно узнаем при загрузке
    pick = DataValidation(type="list", formula1=f"='{ROLES}'!$A$2:$A${len(roles) + 1}", allow_blank=True)
    people.add_data_validation(pick)
    pick.add("A2:A200")
    # роль, которой нет в списке, пришла текстом от пользователя: формулой она стать не должна
    defuse(book)

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _style(sheet, widths: list[int]) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, width in enumerate(widths):
        sheet.column_dimensions[chr(ord("A") + index)].width = width
    sheet.freeze_panes = "A2"


def parse(
    content: bytes, name: str, facility_id: str, operation_ids: list[str], current: dict[str, float] | None = None
) -> DataImportResult:
    """current: что человек уже ввел руками. Смены из файла складываются с ними, а не с умолчаниями."""
    if len(content) > MAX_BYTES:
        raise TooLarge(name)
    tables = _read(content, name)
    fields = _client_fields(facility_id, operation_ids)
    roles = staff_service.roles()

    rows: list[DataImportRow] = []
    overrides: dict[str, float] = {}
    staff: list[StaffLine] | None = None
    for sheet, table in tables:
        head = [_norm(cell) for cell in (table[0] if table else [])]
        if any("роль" in cell for cell in head):
            staff = (staff or []) + _staff(sheet, table, roles, rows)
        elif head:
            _object(sheet, table, fields, overrides, rows)

    if staff == []:
        # Лист штата есть, а исправных строк нет: штат человека не трогаем, а не стираем
        staff = None
        rows.append(
            DataImportRow(
                sheet=STAFF,
                row=1,
                field="",
                value="",
                status="warn",
                message="Исправных строк штата нет, оставили штат, который был",
            )
        )
    _impossible(overrides, rows, current or {})
    return DataImportResult(
        overrides=overrides,
        staff=staff,
        rows=rows,
        applied=len(overrides) + len(staff or []),
        # предупреждение про пустое поле ничего не подставляет: в "подставим с замечанием" его не считаем
        warnings=sum(1 for row in rows if row.status == "warn" and KEPT not in row.message),
        errors=sum(1 for row in rows if row.status == "error"),
    )


def _read(content: bytes, name: str) -> list[tuple[str, list[list[str]]]]:
    lower = name.lower()
    if lower.endswith(".xlsx"):
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as error:  # битый файл, не xlsx внутри
            raise WrongFile("Excel не открылся: файл поврежден или это не .xlsx") from error
        tables = []
        for sheet in book.worksheets:
            if sheet.title == ROLES:
                continue
            # Область листа файл объявляет сам: 5 КБ с областью A1:XFD2000 дают 32 млн пустых ячеек.
            # Поэтому область забываем и читаем не больше, чем бывает в шаблоне с запасом
            sheet.reset_dimensions()
            rows = sheet.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLUMNS, values_only=True)
            table = [["" if cell is None else str(cell) for cell in row] for row in rows]
            tables.append((sheet.title, table))
        return tables
    if lower.endswith(".csv"):
        text = _decode(content)
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=";,\t") if text.strip() else csv.excel
        except csv.Error:
            dialect = Semicolon  # разделитель не угадался: читаем как шаблон, дальше скажем, каких колонок нет
        rows = itertools.islice(csv.reader(io.StringIO(text), dialect), MAX_ROWS)
        return [("CSV", [row[:MAX_COLUMNS] for row in rows])]
    raise WrongFile("Нужен файл Excel (.xlsx) или CSV")


def _decode(content: bytes) -> str:
    # Excel по-русски сохраняет CSV в cp1251, остальные в UTF-8
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise WrongFile("CSV не прочитался: сохраните его в UTF-8")


def _object(
    sheet: str, table: list[list[str]], fields: list[Parameter], overrides: dict[str, float], rows: list[DataImportRow]
) -> None:
    head = [_norm(cell) for cell in table[0]]
    code, label, value, unit = (_column(head, words) for words in (["код"], ["параметр"], ["значение"], ["единица"]))
    if value is None or (code is None and label is None):
        rows.append(
            DataImportRow(
                sheet=sheet,
                row=1,
                field="",
                value="",
                status="error",
                message="Не нашли колонки «Параметр» и «Значение». Возьмите шаблон",
            )
        )
        return
    by_path = {field.path: field for field in fields}
    by_label = {_norm(field.label): field for field in fields}
    # файлы по старому шаблону: поле переименовали, прежнее название тоже узнаем
    for old, new in OLD_LABELS.items():
        if _norm(new) in by_label:
            by_label.setdefault(_norm(old), by_label[_norm(new)])
    seen: set[str] = set()
    for number, cells in enumerate(table[1:], start=2):
        raw = _cell(cells, value)
        key = _cell(cells, code) or _cell(cells, label)
        if not key and not raw:
            continue
        field = by_path.get(_cell(cells, code)) or by_label.get(_norm(_cell(cells, label)))
        shown = field.label if field else key

        def mark(status: str, message: str = "", shown=shown, raw=raw, number=number) -> None:
            rows.append(DataImportRow(sheet=sheet, row=number, field=shown, value=raw, status=status, message=message))

        if field is None:
            mark("error", "Такого параметра нет в расчете для выбранных задач")
            continue
        seen.add(field.path)
        if not raw:
            if field.key:
                mark(
                    "error",
                    f"Обязательное поле пустое, оставили значение по умолчанию: {_default(field)}. Впишите свое",
                )
            else:
                mark("warn", f"Пусто. {KEPT}: {_default(field)}")
            continue
        amount = _number(raw)
        if amount is None:
            mark("error", "Не число")
            continue
        given = _cell(cells, unit)
        if given and _unit(given) != _unit(field.unit):
            mark("error", f"Ждем в «{field.unit}», в файле «{given}»: переведите в нашу единицу")
            continue
        if amount < 0:
            mark("error", "Меньше нуля быть не может")
            continue
        if field.path in overrides:
            mark("warn", "Параметр встречается дважды, взяли последнее значение")
        elif amount < field.min or amount > field.max:
            mark(
                "warn",
                f"Вне того, что мы видели в данных: от {field.min:g} до {field.max:g} {field.unit}. "
                "Подставим, а расчет пометим",
            )
        else:
            mark("ok")
        overrides[field.path] = amount

    # Главные поля без строки в файле: молча подставлять значение по умолчанию нельзя,
    # человек решит, что расчет стоит на его цифрах
    for field in fields:
        if field.key and field.path not in seen:
            rows.append(
                DataImportRow(
                    sheet=sheet,
                    row=0,
                    field=field.label,
                    value="",
                    status="warn",
                    message=f"Обязательного поля нет в файле. {KEPT}: {_default(field)}",
                )
            )


def _default(field: Parameter) -> str:
    """Значение по умолчанию так, как его видно на экране: 10 000 м2."""
    # пробелы неразрывные: в отчете узкая колонка, и "2 000 паллет" не должно рваться на строки
    value = f"{field.value:,.10g}".replace(",", "\u00a0").replace(".", ",")
    return f"{value}\u00a0{field.unit}".strip()


def _staff(sheet: str, table: list[list[str]], roles: list[Role], rows: list[DataImportRow]) -> list[StaffLine]:
    head = [_norm(cell) for cell in table[0]]
    role_col, seats, busy, salary, outside = (
        _column(head, words) for words in (["роль"], ["мест"], ["занято"], ["оклад"], ["подрядчик"])
    )
    known = {_norm(role.name): role for role in roles} | {_norm(role.id): role for role in roles}
    lines: list[StaffLine] = []
    for number, cells in enumerate(table[1:], start=2):
        name = _cell(cells, role_col)
        if not name and not _cell(cells, seats):
            continue

        def mark(status: str, message: str = "", name=name, number=number) -> None:
            rows.append(DataImportRow(sheet=sheet, row=number, field=name, value="", status=status, message=message))

        role = known.get(_norm(name))
        if role is None:
            mark("error", "Такой роли нет в справочнике: выберите из листа «Роли»")
            continue
        headcount = _number(_cell(cells, seats))
        if headcount is None or headcount < 0:
            mark("error", "Мест в штате: нужно число не меньше нуля")
            continue
        filled = _number(_cell(cells, busy))
        filled = headcount if filled is None else filled
        pay = _number(_cell(cells, salary))
        flag = _norm(_cell(cells, outside))
        if flag not in YES | NO:
            mark("error", "Люди подрядчика: напишите «да» или «нет»")
            continue
        notes = []
        if pay is None:
            pay = role.salary_month
            notes.append(f"оклад не указан, взяли рыночный {pay:,.0f} руб".replace(",", " "))
        if filled > headcount:
            notes.append("занято больше, чем мест в штате")
        mark("warn" if notes else "ok", "; ".join(notes).capitalize())
        lines.append(
            StaffLine(
                id=f"file-{number}",
                role=role.id,
                headcount=headcount,
                filled=filled,
                salary_month=pay,
                contractor=flag in YES,
            )
        )
    return lines


def _impossible(overrides: dict[str, float], rows: list[DataImportRow], current: dict[str, float]) -> None:
    """Режим работы проверяем так же, как при ручном вводе: больше 24 часов в сутки не подставляем.
    Смены из файла складываем с тем, что человек уже ввел: 13 часов при его одной смене это 13 часов в сутки."""
    mine = {path: value for path, value in current.items() if ".schedule." in path and 0 <= value < math.inf}
    try:
        try:
            calculation.model_with_overrides({**mine, **overrides})
        except calculation.UnknownPath:
            calculation.model_with_overrides(overrides)  # правки с другого объекта: судим только по файлу
    except calculation.Impossible as error:
        for path in [path for path in overrides if ".schedule." in path]:
            del overrides[path]
        for row in rows:
            # строки без номера это поля, которых нет в файле: их значение не из файла, ошибкой не помечаем
            if row.row and row.status != "error" and ("смен" in row.field.lower() or "дней" in row.field.lower()):
                row.status = "error"
                row.message = str(error)


def _column(head: list[str], words: list[str]) -> int | None:
    return next((index for index, cell in enumerate(head) if any(word in cell for word in words)), None)


def _cell(cells: list[str], index: int | None) -> str:
    if index is None or index >= len(cells):
        return ""
    return str(cells[index]).strip()


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", str(text).strip().lower().replace("ё", "е"))


def _unit(text: str) -> str:
    """Единица для сравнения: в шаблоне раньше писали "руб/кВт*ч", теперь "₽/кВт·ч", старый файл тоже подходит"""
    return _norm(text).replace("руб.", "₽").replace("руб", "₽").replace("*", "·")


def _number(raw: str) -> float | None:
    cleaned = re.sub(r"[\s\u00a0]", "", str(raw)).replace(",", ".")
    if not cleaned:
        return None
    try:
        value = float(cleaned)
    except ValueError:
        return None
    return value if math.isfinite(value) else None  # float() понимает и слова nan и inf
