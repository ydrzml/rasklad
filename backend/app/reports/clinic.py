"""Отчет медучреждения: PDF и Excel по тому же итогу, что шаг итога (POST /api/clinic/economics).

Разделы по смыслу те же, что у склада: ответ, три сценария, затраты по годам, парк и смена на
корпусе, что с людьми, что может сдвинуть вывод и откуда цифры. Цифры не пересчитываем: берем
ответ сервиса медучреждения как есть, отчет и экран не расходятся. Оформление и шрифты общие
с отчетом склада (pdf.py, excel.py).
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.graphics.shapes import Circle, Drawing, Line, PolyLine, Rect, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import BaseDocTemplate, Frame, Indenter, PageTemplate, Paragraph, Spacer, Table

from app.reports import excel, pdf
from app.reports.cells import defuse
from app.reports.content import MARK, MOSCOW, Saved, grouped, mln, percent, term
from app.schemas.clinic import MoneyRequest, MoneyResult, MoneyScenario
from app.services import clinic_money
from app.services.versions import model_version

FACILITY = "Медучреждение"
NAMES = {"baseline": "Без роботов", "purchase": "Покупка", "raas": "Аренда", "loan": "Покупка в кредит"}
TONES = {**pdf.TONES, "loan": colors.HexColor("#5a3ea8")}
SOURCE_GROUPS = {
    "Объект": "Ваш объект",
    "Роботы": "Роботы",
    "Смена": "Смена: лифт, двери, сроки, люди",
    "Деньги": "Деньги",
}


@dataclass
class ClinicReport:
    created: datetime
    request: MoneyRequest
    result: MoneyResult
    model_version: str
    saved: Saved | None = None

    @property
    def work(self) -> str:
        return ", ".join(task.lower() for task in self.result.tasks) or "задачи не выбраны"

    @property
    def fleet(self) -> str:
        return " и ".join(f"{one.count} {one.name}" for one in self.result.fleets if one.count) or "без роботов"

    @property
    def file_stem(self) -> str:
        version = f"-v{self.saved.number}" if self.saved else ""
        return f"otsenka-clinic{version}-{self.created:%Y-%m-%d}"

    @property
    def stamp(self) -> str:
        version = f"версия {self.saved.number} · " if self.saved else ""
        return f"{version}расчет от {self.created:%d.%m.%Y %H:%M} мск · модель {self.model_version}"

    def scenario(self, scenario_id: str) -> MoneyScenario | None:
        return next((one for one in self.result.scenarios if one.id == scenario_id), None)

    @property
    def shown(self) -> list[MoneyScenario]:
        """Сценарии в отчете: без роботов, покупка и аренда всегда, кредит только когда его добавили"""
        return [one for one in self.result.scenarios if one.id != "loan" or one.available]


def build(request: MoneyRequest) -> ClinicReport:
    """Отчет по свежему расчету: тот же запрос, что ушел на шаг итога"""
    return ClinicReport(datetime.now(MOSCOW), request, clinic_money.estimate(request), model_version())


def build_saved(request: MoneyRequest, result: MoneyResult, saved: Saved, version: str) -> ClinicReport:
    """Отчет по сохраненной версии: цифры из сохранения, без пересчета"""
    return ClinicReport(saved.saved_at.astimezone(MOSCOW), request, result, version, saved)


# --- слова ---------------------------------------------------------------------------------


def answer(report: ClinicReport) -> tuple[str, str]:
    """Вывод заголовком и строкой под ним, по тем же правилам, что заголовок шага итога (steps/clinic.ts)"""
    result, horizon = report.result, report.result.horizon_years
    baseline, buy, rent = report.scenario("baseline"), report.scenario("purchase"), report.scenario("raas")
    if not baseline or not buy or not rent:
        return "Итог не посчитан", "В ответе нет сценариев для сравнения."
    service = result.shift.service
    ready = (service.completion_share is None or service.completion_share >= 0.999) and (
        service.cleaning_coverage is None or service.cleaning_coverage >= 0.999
    )
    if result.shift.search is not None and result.shift.search.needed is None:
        ready = False
    if not ready:
        return (
            "Достаточность парка не подтверждена",
            "Цель доставки или уборки не достигнута. Денежный эффект относится только к выполненному объему; "
            "сначала устраните ограничения и повторите прогон.",
        )
    if result.assumptions.get("cash_realization_share", 1) == 0:
        return (
            "Сокращение бюджета не подтверждено",
            "Роботы могут высвободить время для других задач. Зарплата остается в затратах: стоимость этого "
            "времени не считается денежной экономией.",
        )

    def pays(one: MoneyScenario) -> bool:
        gain = baseline.tco_rub - one.tco_rub
        return one.available and gain > 0.01 and one.payback_years is not None and 0 <= one.payback_years <= horizon

    if not pays(buy) and not pays(rent):
        return (
            "Окупаемость с положительным результатом не подтверждена",
            f"На текущих входных данных расходы не покрываются за {horizon} лет. Что держит вывод, в разделах "
            "о людях и о том, что может его сдвинуть.",
        )
    if pays(buy) and (not pays(rent) or buy.tco_rub <= rent.tco_rub):
        return (
            f"Выгоднее покупка: окупится за {term(buy.payback_years)}",
            f"За {horizon} лет дешевле, чем без роботов, на {mln(baseline.tco_rub - buy.tco_rub)} млн ₽. "
            f"Вложить {mln(buy.start_rub)} млн ₽, экономия {mln(buy.effect_year1_rub)} млн ₽ в первый год.",
        )
    return (
        "Выгоднее аренда",
        f"За {horizon} лет дешевле, чем без роботов, на {mln(baseline.tco_rub - rent.tco_rub)} млн ₽.",
    )


def payback(one: MoneyScenario, horizon: int) -> str:
    if one.id == "baseline":
        return "-"
    return f"не за {horizon} лет" if one.payback_years is None else term(one.payback_years)


def cumulative(one: MoneyScenario) -> list[float]:
    """Сколько всего потрачено к концу каждого года: старт и расходы лет нарастающим итогом"""
    points, total = [one.start_rub], one.start_rub
    for value in one.years_rub:
        total += value
        points.append(total)
    return points


def shift_rows(report: ClinicReport) -> list[tuple[str, str]]:
    """Смена на корпусе: что прогнали и что вышло, теми же словами, что на экране"""
    shift = report.result.shift
    rows = [
        ("Корпус", f"этажей {shift.floors}, постов {shift.wards}, лифтов {shift.lifts}"),
        ("Окно прогона", f"{grouped(shift.hours, 0)} ч с {shift.shift_start_h:02d}:00"),
        ("Парк", report.fleet),
    ]
    if shift.jobs:
        rows += [
            ("Заявок выполнено", f"{shift.done} из {shift.jobs}"),
            ("Вовремя", "нет заявок" if shift.on_time_share is None else percent(shift.on_time_share)),
            (
                "В среднем от заявки до поста",
                "нет" if shift.avg_delivery_min is None else f"{grouped(shift.avg_delivery_min, 0)} мин",
            ),
        ]
    if shift.lift_busy_share is not None:
        rows.append(("Самый занятый лифт", f"{percent(shift.lift_busy_share)} смены"))
    if shift.cleaning_required_m2:
        rows.append(
            ("Вымыто коридоров", f"{grouped(shift.cleaned_m2, 0)} из {grouped(shift.cleaning_required_m2, 0)} м²")
        )
    rows.append(("Узкое место", shift.bottleneck))
    return rows


def people_rows(report: ClinicReport) -> list[tuple[str, float, str]]:
    people = report.result.people
    return [
        ("Ручной эквивалент выполненной работы", people.manual_hours_per_shift, "ч в смену"),
        ("Остается людям с роботами, с помощью роботам", people.remaining_hours_per_shift, "ч в смену"),
        ("Чистое высвобождение", people.hours_per_shift, "ч в смену"),
        ("То же за год", people.hours_per_year, "ч в год"),
        ("Эквивалент ставок (мера времени, не увольнения)", people.fte, "ставок"),
        ("Стоимость чистого времени", people.released_value_year_rub, "₽ в год"),
        ("Эффект фонда оплаты труда по сценарию", people.cash_saving_year_rub, "₽ в год"),
    ]


def limits(report: ClinicReport) -> list[str]:
    """Что может сдвинуть вывод: что требует уточнения, предупреждения расчета и проверки корпуса"""
    found: list[str] = []
    checks = [one.text for one in report.result.shift.checks if one.level in ("warning", "error")]
    for line in [*report.result.readiness, *report.result.warnings, *checks]:
        if line not in found:
            found.append(line)
    return found


# --- PDF ----------------------------------------------------------------------------------


def render_pdf(report: ClinicReport) -> bytes:
    pdf._register()
    s = pdf._styles()
    out = BytesIO()
    x0, y0, x1, y1 = pdf.FRAME
    inner = x1 - x0
    body = Frame(x0, y0 + 10 * mm, inner, y1 - y0 - 22 * mm, 0, 0, 0, 0)
    doc = BaseDocTemplate(out, pagesize=A4, title=f"Оценка роботизации: медучреждение, {report.work}", author="Расклад")
    doc.addPageTemplates([PageTemplate("sheet", [body], onPage=lambda c, d: _sheet(c, report))])
    main = inner - pdf.LABEL - pdf.GAP
    count = iter(range(1, 100))

    def no() -> str:
        return f"{next(count):02d}"

    head, sub = answer(report)
    horizon = report.result.horizon_years
    story: list = []
    story += pdf._section(
        "",
        "Объект",
        None,
        [
            Paragraph(f"Медучреждение,<br/>{escape(report.work)}", s["title"]),
            Spacer(0, 6),
            Paragraph(escape(report.fleet), s["soft"]),
            Spacer(0, 6),
            Paragraph(escape(_version(report)), s["note"]),
        ],
        s,
        first_block=True,
    )
    story += pdf._section(
        no(),
        "Ответ",
        None,
        [Paragraph(escape(head), s["answer"]), Spacer(0, 4), Paragraph(escape(sub), s["base"])],
        s,
    )
    story += pdf._section(
        no(),
        "Три сценария",
        f"Стоимость владения: вложения на старте и все расходы за {horizon} лет, включая сохраняющуюся "
        "зарплату. NPV учитывает время поступления денег.",
        [_scenarios(report, main)],
        s,
    )
    story += pdf._section(
        no(),
        "Затраты по годам",
        "Сколько всего потрачено к концу каждого года, млн ₽. Где линия с роботами уходит ниже линии "
        "без роботов, вложения окупились.",
        [_chart(report, main), Spacer(0, 8), _years(report, main)],
        s,
    )
    story += pdf._section(
        no(),
        "Парк и смена на корпусе",
        "Выбранные роботы на корпусе с шага плана: двери, очередь к лифту, зарядка. Парк доставки "
        "самый маленький, при котором вовремя не меньше цели.",
        [_pairs(shift_rows(report), main), Spacer(0, 8), _fleets(report, s, main)],
        s,
    )
    story += pdf._section(
        no(),
        "Что с людьми",
        "Часы санитаров и уборщиц, которые забирают роботы, после помощи роботам.",
        [_people(report, main)],
        s,
    )
    risks = limits(report)
    if risks:
        story += pdf._section(
            no(),
            "Что может сдвинуть вывод",
            None,
            [pdf._numbered([Paragraph(escape(line), s["base"]) for line in risks], main)],
            s,
        )
    if report.result.not_counted:
        story += pdf._section(
            no(),
            "Чего нет в деньгах",
            None,
            [pdf._numbered([Paragraph(escape(line), s["base"]) for line in report.result.not_counted], main)],
            s,
        )
    sources = report.result.sources
    story += pdf._section(
        no(),
        "Откуда цифры",
        f"Значений: {len(sources)}, без источника (оценка F): {sum(1 for one in sources if one.trust == 'F')}. "
        "S закон или норматив, A проверено, B два источника, C один источник, D каталог или СМИ, "
        "E источники расходятся, F наше допущение.",
        [_sources(report, s, main)],
        s,
    )
    story += [
        Spacer(0, 14),
        Indenter(left=pdf.LABEL + pdf.GAP),
        Paragraph(
            'Это предварительная оценка для решения "считать дальше или нет", а не коммерческое предложение. '
            "Расчет не подтверждает санитарную пригодность оборудования: для проекта нужны наблюдаемые смены "
            "на объекте и предложения поставщиков.",
            s["note"],
        ),
        Indenter(left=-(pdf.LABEL + pdf.GAP)),
    ]
    doc.build(story, canvasmaker=_Numbered)
    return out.getvalue()


def _version(report: ClinicReport) -> str:
    lines = []
    if report.saved:
        lines.append(f'Версия {report.saved.number} проекта "{report.saved.project}", цифры из сохранения.')
    lines.append(f"Расчет от {report.created:%d.%m.%Y %H:%M} мск, модель {report.model_version}.")
    lines.append(f"Горизонт {report.result.horizon_years} лет. {MARK}.")
    return " ".join(lines)


def _scenarios(report: ClinicReport, width: float) -> Table:
    horizon = report.result.horizon_years
    baseline = report.scenario("baseline")
    rows: list[list] = [["Сценарий", "На старте", f"За {horizon} лет", "Меньше, чем без роботов", "Окупится", "NPV"]]
    for one in report.shown:
        if not one.available:
            rows.append([f"{NAMES[one.id]}: тариф уточняется", "-", "-", "-", "-", "-"])
            continue
        gain = baseline.tco_rub - one.tco_rub if baseline and one.id != "baseline" else None
        rows.append(
            [
                NAMES[one.id],
                f"{mln(one.start_rub)} млн",
                f"{mln(one.tco_rub)} млн",
                "-" if gain is None else f"{mln(gain)} млн",
                Paragraph(escape(payback(one, horizon)), _right_cell()),
                "-" if one.id == "baseline" else f"{mln(one.npv_rub)} млн",
            ]
        )
    soft = pdf._styles()["cellsoft"]
    rows[0] = [rows[0][0], *[Paragraph(escape(name), _right_cell(soft)) for name in rows[0][1:]]]
    return pdf._table(rows, [width * share for share in (0.2, 0.13, 0.13, 0.18, 0.2, 0.16)])


def _right_cell(base=None):
    """Ячейка с переносом строки, прижатая вправо, как числа таблицы"""
    style = base or pdf._styles()["cell"]
    return ParagraphStyle(f"{style.name}-right", parent=style, alignment=2)


def _years(report: ClinicReport, width: float) -> Table:
    shown = [one for one in report.shown if one.available]
    lines = {one.id: cumulative(one) for one in shown}
    rows: list[list] = [["Год", *[NAMES[one.id] for one in shown]]]
    for year in range(max((len(points) for points in lines.values()), default=0)):
        rows.append(["старт" if year == 0 else f"{year} год", *[f"{mln(lines[one.id][year])} млн" for one in shown]])
    first = width * 0.2
    return pdf._table(rows, [first, *[(width - first) / max(1, len(shown))] * len(shown)])


def _chart(report: ClinicReport, width: float) -> Drawing:
    """Накопленные затраты. На черно-белой печати сценарии различаются линией и меткой"""
    height = 150
    pad_l, pad_r, pad_t, pad_b = 26, 104, 6, 18
    series = [(one, cumulative(one)) for one in report.shown if one.available]
    years = max(1, max(len(points) for _, points in series) - 1)
    top = pdf._nice(max(max(points) for _, points in series) / 1_000_000)
    drawing = Drawing(width, height)

    def x(year: float) -> float:
        return pad_l + year / years * (width - pad_l - pad_r)

    def y(value: float) -> float:
        return pad_b + value / 1_000_000 / top * (height - pad_t - pad_b)

    for tick in (0, top / 2, top):
        at = y(tick * 1e6)
        base = tick == 0
        tone, thick = (pdf.INK, 0.75) if base else (pdf.HAIR, 0.5)
        drawing.add(Line(pad_l, at, width - pad_r, at, strokeColor=tone, strokeWidth=thick))
        label = grouped(tick, 1 if tick < 10 else 0)
        drawing.add(
            String(pad_l - 5, at - 2.5, label, fontName="Mono", fontSize=7, fillColor=pdf.SOFT, textAnchor="end")
        )
    for year in range(years + 1):
        label = "старт" if year == 0 else f"{year} год"
        drawing.add(String(x(year), 5, label, fontName="Text", fontSize=7, fillColor=pdf.SOFT, textAnchor="middle"))
    placed: list[float] = []
    for one, points in sorted(series, key=lambda item: item[1][-1]):
        tone = TONES.get(one.id, pdf.SOFT)
        line = PolyLine([c for n, v in enumerate(points) for c in (x(n), y(v))], strokeColor=tone, strokeWidth=1.4)
        if one.id == "baseline":
            line.strokeDashArray = [4, 2.5]
        drawing.add(line)
        for n, value in enumerate(points):
            if one.id in ("purchase", "loan"):
                drawing.add(
                    Rect(
                        x(n) - 2.2, y(value) - 2.2, 4.4, 4.4, fillColor=tone, strokeColor=colors.white, strokeWidth=0.6
                    )
                )
            elif one.id == "raas":
                drawing.add(Circle(x(n), y(value), 2.4, fillColor=colors.white, strokeColor=tone, strokeWidth=1.1))
        end = y(points[-1])
        if placed and end - placed[-1] < 11:
            end = placed[-1] + 11
        placed.append(end)
        name = NAMES[one.id]
        drawing.add(String(width - pad_r + 8, end - 3, name, fontName="Text", fontSize=8, fillColor=pdf.INK))
        shift = pdfmetrics.stringWidth(name + " ", "Text", 8)
        drawing.add(
            String(width - pad_r + 8 + shift, end - 3, mln(points[-1]), fontName="Mono", fontSize=8, fillColor=pdf.SOFT)
        )
    return drawing


def _pairs(rows: list[tuple[str, str]], width: float) -> Table:
    style = pdf._styles()["cell"]
    cells = [[name, Paragraph(escape(value), style)] for name, value in rows]
    return pdf._table(cells, [width * 0.4, width * 0.6], None, head=False)


def _fleets(report: ClinicReport, s: dict, width: float) -> Table:
    rows: list[list] = [["Задача", "Робот", "Шт.", "Цена за шт.", "Откуда цена"]]
    for one in report.result.fleets:
        rows.append(
            [
                Paragraph(escape(one.task), s["cell"]),
                Paragraph(escape(one.name), s["cell"]),
                str(one.count),
                f"{mln(one.price_rub)} млн",
                Paragraph(escape(one.price_source), s["cellsoft"]),
            ]
        )
    return pdf._table(rows, [width * 0.18, width * 0.2, width * 0.08, width * 0.14, width * 0.4], 2)


def _people(report: ClinicReport, width: float) -> Table:
    rows: list[list] = [["Показатель", "Значение"]]
    for name, value, unit in people_rows(report):
        rows.append([name, f"{grouped(value, 0)} {unit}" if unit.startswith("₽") else f"{grouped(value, 1)} {unit}"])
    return pdf._table(rows, [width * 0.65, width * 0.35])


def _sources(report: ClinicReport, s: dict, width: float) -> Table:
    rows: list[list] = [["Значение", "Сколько", "Источник", "Оценка"]]
    for group, title in SOURCE_GROUPS.items():
        lines = [one for one in report.result.sources if one.group == group]
        if not lines:
            continue
        rows.append([Paragraph(f"<b>{escape(title)}</b>", s["cell"]), "", "", ""])
        for one in lines:
            when = f" · {one.date}" if one.date else ""
            rows.append(
                [
                    Paragraph(escape(one.name), s["cell"]),
                    Paragraph(escape(one.value), s["cell"]),
                    Paragraph(escape(one.source + when), s["cellsoft"]),
                    one.trust,
                ]
            )
    return pdf._table(rows, [width * 0.27, width * 0.15, width * 0.5, width * 0.08], 3)


class _Numbered(Canvas):
    """Холст, который знает число листов: в нижней строке "2 / 9"."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._pages: list[dict] = []

    def showPage(self):  # noqa: N802 - имя метода задает reportlab
        self._pages.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._pages)
        for page in self._pages:
            self.__dict__.update(page)
            _footer(self, self._clinic_report, total)
            super().showPage()
        super().save()


def _sheet(canvas: Canvas, report: ClinicReport) -> None:
    x0, _, x1, y1 = pdf.FRAME
    canvas.saveState()
    canvas.setStrokeColor(pdf.LINE)
    canvas.setLineWidth(0.6)
    canvas.line(x0, y1 - 2.5 * mm, x1, y1 - 2.5 * mm)
    pdf._caps(canvas, "ОЦЕНКА РОБОТИЗАЦИИ · МЕДУЧРЕЖДЕНИЕ", x0, y1, 7, pdf.SOFT)
    right = "ПРЕДВАРИТЕЛЬНАЯ ОЦЕНКА"
    pdf._caps(canvas, right, x1 - pdf._caps_width(right, 7), y1, 7, pdf.BLUE)
    canvas.restoreState()
    canvas._clinic_report = report


def _footer(canvas: Canvas, report: ClinicReport, total: int) -> None:
    x0, y0, x1, _ = pdf.FRAME
    canvas.saveState()
    canvas.setStrokeColor(pdf.HAIR)
    canvas.setLineWidth(0.5)
    canvas.line(x0, y0 + 4 * mm, x1, y0 + 4 * mm)
    sheet = f"{canvas.getPageNumber()} / {total}"
    right = x1 - pdfmetrics.stringWidth(sheet, "Mono", 7.5)
    left = f"{FACILITY} · {report.fleet} · {report.stamp}"
    canvas.setFont("Text", 7.5)
    canvas.setFillColor(pdf.SOFT)
    canvas.drawString(x0, y0, pdf._fit(left, "Text", 7.5, right - x0 - 8 * mm))
    canvas.setFont("Mono", 7.5)
    canvas.setFillColor(pdf.INK)
    canvas.drawString(right, y0, sheet)
    canvas.restoreState()


# --- Excel --------------------------------------------------------------------------------


def render_xlsx(report: ClinicReport) -> bytes:
    """Листы: итог, по годам, парк и смена, люди, ограничения, источники. Числа лежат числами"""
    book = Workbook()
    totals = _x_years(book.active, report)
    _x_summary(book.create_sheet("Итог", 0), report, totals)
    _x_shift(book.create_sheet("Парк и смена"), report)
    _x_people(book.create_sheet("Люди"), report)
    _x_limits(book.create_sheet("Ограничения"), report)
    _x_sources(book.create_sheet("Источники"), report)
    book.active = 0
    defuse(book)
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def _x_sheet(sheet: Worksheet, title: str, report: ClinicReport, widths: list[int]) -> int:
    """Шапка листа как у склада: что это, по какому объекту, когда и какой моделью посчитано"""
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = title
    sheet["A1"].font = excel.TITLE
    sheet["A2"] = f"{FACILITY}, {report.work} · {report.fleet} · {report.stamp}"
    sheet["A2"].font = excel.MUTED
    sheet["A3"] = MARK
    sheet["A3"].font = excel.MARK
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "5:5"
    sheet.oddFooter.left.text = "Расклад · предварительная оценка"
    sheet.oddFooter.right.text = "Лист &P из &N"
    return 5


def _x_years(sheet: Worksheet, report: ClinicReport) -> dict[str, str]:
    """По годам: расходы года и нарастающий итог формулой. Отдает адрес итога сценария для листа "Итог"."""
    sheet.title = "По годам"
    shown = [one for one in report.shown if one.available]
    row = _x_sheet(sheet, "Затраты по годам, ₽", report, [16, *[18] * (2 * len(shown))])
    titles = ["Год"]
    for one in shown:
        titles += [f"{NAMES[one.id]}: за год", f"{NAMES[one.id]}: всего"]
    excel._head(sheet, row, titles)
    first = row + 1
    length = max((len(one.years_rub) for one in shown), default=0)
    for year in range(length + 1):
        values: list = ["старт" if year == 0 else f"{year} год"]
        for index, one in enumerate(shown):
            spend = one.start_rub if year == 0 else one.years_rub[year - 1]
            column = get_column_letter(2 + 2 * index)
            here = f"{column}{first + year}"
            values += [spend, f"={here}" if year == 0 else f"={here}+{_right(column)}{first + year - 1}"]
        excel._put(sheet, first + year, values, excel.RUB)
    return {
        one.id: f"'По годам'!{_right(get_column_letter(2 + 2 * index))}{first + length}"
        for index, one in enumerate(shown)
    }


def _right(column: str) -> str:
    return get_column_letter(column_index_from_string(column) + 1)


def _x_summary(sheet: Worksheet, report: ClinicReport, totals: dict[str, str]) -> None:
    head, sub = answer(report)
    horizon = report.result.horizon_years
    row = _x_sheet(sheet, "Итог", report, [26, 18, 20, 22, 18, 18])
    excel._head(
        sheet, row, ["Сценарий", "На старте, ₽", f"За {horizon} лет, ₽", "Дешевле без роботов, ₽", "Окупится", "NPV, ₽"]
    )
    baseline_row = None
    for one in report.shown:
        row += 1
        if not one.available:
            excel._put(sheet, row, [NAMES[one.id], "-", "уточняется", "-", "-", "-"])
            continue
        if one.id == "baseline":
            baseline_row = row
        total = f"={totals[one.id]}" if one.id in totals else one.tco_rub
        gain = "-" if one.id == "baseline" or baseline_row is None else f"=C{baseline_row}-C{row}"
        npv = "-" if one.id == "baseline" else one.npv_rub
        excel._put(sheet, row, [NAMES[one.id], one.start_rub, total, gain, payback(one, horizon), npv], excel.RUB)
    row += 2
    sheet.cell(row=row, column=1, value=head).font = excel.STRONG
    note = sheet.cell(row=row + 1, column=1, value=sub)
    note.font = excel.TEXT
    note.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=6)
    sheet.row_dimensions[row + 1].height = 45


def _x_shift(sheet: Worksheet, report: ClinicReport) -> None:
    row = _x_sheet(sheet, "Парк и смена на корпусе", report, [34, 30, 10, 18, 60])
    excel._head(sheet, row, ["Показатель", "Значение"])
    for name, value in shift_rows(report):
        row += 1
        excel._put(sheet, row, [name, value])
    row += 2
    excel._head(sheet, row, ["Задача", "Робот", "Штук", "Цена за штуку, ₽", "Откуда цена"])
    for one in report.result.fleets:
        row += 1
        excel._put(sheet, row, [one.task, one.name, one.count, one.price_rub, one.price_source], excel.RUB)
        sheet.cell(row=row, column=3).number_format = "0"


def _x_people(sheet: Worksheet, report: ClinicReport) -> None:
    row = _x_sheet(sheet, "Что с людьми", report, [52, 18, 14])
    excel._head(sheet, row, ["Показатель", "Значение", "Единица"])
    for name, value, unit in people_rows(report):
        row += 1
        excel._put(sheet, row, [name, round(value, 2), unit], excel.RUB if unit.startswith("₽") else excel.ONE)


def _x_limits(sheet: Worksheet, report: ClinicReport) -> None:
    row = _x_sheet(sheet, "Что может сдвинуть вывод и чего нет в деньгах", report, [110])
    excel._head(sheet, row, ["Что может сдвинуть вывод"])
    for line in limits(report):
        row += 1
        excel._put(sheet, row, [line])
    row += 2
    excel._head(sheet, row, ["Чего нет в деньгах"])
    for line in report.result.not_counted:
        row += 1
        excel._put(sheet, row, [line])


def _x_sources(sheet: Worksheet, report: ClinicReport) -> None:
    row = _x_sheet(sheet, "Откуда цифры", report, [18, 40, 22, 70, 10, 8])
    excel._head(sheet, row, ["Раздел", "Значение", "Сколько", "Источник", "Дата", "Оценка"])
    for one in report.result.sources:
        row += 1
        group = SOURCE_GROUPS.get(one.group, one.group)
        excel._put(sheet, row, [group, one.name, one.value, one.source, one.date or "", one.trust])
