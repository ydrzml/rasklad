"""Расчетные таблицы в Excel (ТЗ, п. 3.8). Листы: итог, по годам, затраты, что будет если, парк и штат,
смена на плане, параметры, источники.

Оформление как у отчета: шапка листа вместо штампа, без заливок и без сетки Excel, заголовки колонок
с единицами и линией снизу, шапка закреплена. Числа лежат числами. Главное на листе «Итог» посчитано
формулами из листа «По годам»: видно, откуда берется каждая цифра, и ее можно проверить у себя.
Шрифт Arial: Onest есть не на каждом компьютере, а таблицу откроют где угодно.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.reports import names
from app.reports.cells import defuse
from app.reports.content import (
    CAPEX_NAMES,
    NAMES,
    OPEX_NAMES,
    Report,
    delta,
    financed,
    irr_text,
    project_irr,
    scenario_name,
)
from app.schemas.economics import SensitivityAll

FONT = "Arial"
RUB = '#,##0 "₽";-#,##0 "₽";"0 ₽"'
MLN = '#,##0.0,, "млн ₽"'
ONE = "0.0"
SHARE = "0%"
INK = "0D141B"
SOFT = "55636F"
BLUE = "0B4FA8"
LINE = "1E4F8E"
HAIR = "C9D2DB"

TITLE = Font(name=FONT, size=14, bold=True, color=INK)
TEXT = Font(name=FONT, size=10, color=INK)
MUTED = Font(name=FONT, size=9, color=SOFT)
MARK = Font(name=FONT, size=9, color=BLUE)
HEAD = Font(name=FONT, size=9, bold=True, color=SOFT)
STRONG = Font(name=FONT, size=10, bold=True, color=INK)
TRUST = Font(name=FONT, size=10, bold=True, color=BLUE)
UNDER_HEAD = Border(bottom=Side(style="medium", color=LINE))
UNDER_ROW = Border(bottom=Side(style="thin", color=HAIR))
LAST_ROW = Border(bottom=Side(style="thin", color=INK))


def render(report: Report) -> bytes:
    book = Workbook()
    years = _years(book.active, report)
    _summary(book.create_sheet("Итог", 0), report, years)
    if report.several:
        _tasks(book.create_sheet("По задачам", 1), report)
    _costs(book.create_sheet("Затраты"), report)
    if report.sensitivity:
        _sensitivity(book.create_sheet("Что будет, если"), report)
    _fleet(book.create_sheet("Парк и штат"), report)
    if any(task.shift for task in report.tasks):
        _shift(book.create_sheet("Смена"), report)
    _parameters(book.create_sheet("Параметры"), report)
    _sources(book.create_sheet("Источники"), report)
    book.active = 0
    defuse(book)
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def _sheet(sheet: Worksheet, title: str, report: Report, widths: list[int], landscape: bool = False) -> int:
    """Шапка листа как штамп отчета: что это, по какому объекту, когда и какой моделью посчитано"""
    sheet.sheet_view.showGridLines = False
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet["A1"] = title
    sheet["A1"].font = TITLE
    sheet["A2"] = (
        f"{report.facility.name}, {report.work} · {report.robot_name} · "
        f"расчет от {report.created:%d.%m.%Y %H:%M} мск · модель {report.result.model_version}"
    )
    sheet["A2"].font = MUTED
    sheet["A3"] = "Предварительная оценка, не коммерческое предложение"
    sheet["A3"].font = MARK
    sheet.row_dimensions[1].height = 22
    # печать: по ширине листа, шапка колонок на каждой странице, номер страницы внизу
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "5:5"
    sheet.oddFooter.left.text = "Расклад · предварительная оценка"
    sheet.oddFooter.right.text = "Лист &P из &N"
    return 5


def _head(sheet: Worksheet, row: int, titles: list[str]) -> None:
    for column, name in enumerate(titles, start=1):
        cell = sheet.cell(row=row, column=column, value=name)
        cell.font = HEAD
        cell.border = UNDER_HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="left" if column == 1 else "right")
    sheet.freeze_panes = sheet.cell(row=row + 1, column=2)


def _put(sheet: Worksheet, row: int, values: list, fmt: str | None = None, font: Font = TEXT, border=UNDER_ROW) -> None:
    """Строка таблицы: формат числа ставим числам и формулам, текст оставляем как есть"""
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.font = font
        cell.border = border
        numeric = isinstance(value, int | float) or isinstance(value, str) and value.startswith("=")
        cell.alignment = Alignment(vertical="top", wrap_text=not numeric, horizontal="right" if numeric else None)
        if fmt and column > 1 and numeric:
            cell.number_format = fmt


def _years(sheet: Worksheet, report: Report) -> dict[str, dict[str, str]]:
    """Сценарии по годам. Затраты за год и накопленные затраты посчитаны формулами.
    Возвращает адреса ячеек, на которые ссылается лист «Итог»."""
    sheet.title = "По годам"
    row = _sheet(sheet, "Сценарии по годам", report, [16, 7, 18, 20, 16, 16, 18, 20, 18], landscape=True)
    _head(
        sheet,
        row,
        [
            "Сценарий",
            "Год",
            "Люди на операции, ₽",
            "Роботы: операторы, обслуживание, энергия, аренда, ₽",
            "Вложения, платежи по долгу минус господдержка, ₽",
            "Затраты за год, ₽",
            "Накопленные затраты, ₽",
            "Эффект против работы без роботов, ₽",
            "Денежный поток, ₽",
        ],
    )
    sheet.row_dimensions[row].height = 40
    cells: dict[str, dict[str, str]] = {}
    for s in report.scenarios:
        name = scenario_name(s)
        row += 1
        start = row
        _put(sheet, row, [name, 0, 0, 0, s.investment_year0_rub, f"=C{row}+D{row}+E{row}", f"=F{row}"], RUB)
        for year in s.years:
            row += 1
            values = [
                name,
                year.year,
                year.staff_cost_rub,
                year.opex_total_rub,
                year.investment_rub + (year.financing_rub or 0) - (year.support_rub or 0),
                f"=C{row}+D{row}+E{row}",
                f"=G{row - 1}+F{row}",
                year.effect_rub,
                year.cash_flow_rub,
            ]
            _put(sheet, row, values, RUB)
        for column in range(1, 10):
            sheet.cell(row=row, column=column).border = LAST_ROW
        cells[s.id] = {
            "tco": f"'По годам'!G{row}",
            "invest": f"'По годам'!E{start}",
            "effect": f"'По годам'!H{start + 1}",
        }
    return cells


def _summary(sheet: Worksheet, report: Report, years: dict[str, dict[str, str]]) -> None:
    row = _sheet(sheet, "Три сценария", report, [36, 16, 16, 16, 50])
    scenarios = report.scenarios
    count = len(scenarios)
    _head(sheet, row, ["Показатель", *[scenario_name(s) for s in scenarios], "Как посчитано"])
    sheet.cell(row=row, column=2 + count).alignment = Alignment(horizontal="left", vertical="bottom")
    first = row + 1

    def col(index: int) -> str:
        return get_column_letter(index + 2)

    horizon = report.result.horizon_years
    robots = range(1, count)
    lines: list[tuple[str, list, str, str]] = [
        (
            f"Стоимость владения за {horizon} лет",
            [f"={years[s.id]['tco']}" for s in scenarios],
            MLN,
            "Накопленные затраты последнего года на листе «По годам»",
        ),
        ("Вложения на старте", [f"={years[s.id]['invest']}" for s in scenarios], MLN, "Год 0 на листе «По годам»"),
        (
            f"Затраты за {horizon} лет без вложений",
            [f"={col(i)}{first}-{col(i)}{first + 1}" for i in range(count)],
            MLN,
            "Стоимость владения - вложения на старте",
        ),
        (
            "Дешевле, чем без роботов",
            ["—", *[f"=B{first}-{col(i)}{first}" for i in robots]],
            MLN,
            "Стоимость владения без роботов - стоимость владения сценария",
        ),
        (
            "Годовой эффект, первый год",
            ["—", *[f"={years[s.id]['effect']}" for s in scenarios[1:]]],
            MLN,
            "Год 1 на листе «По годам»",
        ),
        (
            "Окупаемость по ТЗ, лет",
            ["—", *[f"={col(i)}{first + 1}/{col(i)}{first + 4}" for i in robots]],
            ONE,
            "Вложения на старте / годовой эффект",
        ),
        (
            "Окупаемость по потоку, лет",
            ["—", *[s.payback_cumulative_years for s in scenarios[1:]]],
            ONE,
            "Когда накопленная экономия перекроет вложения, с ростом цен и выкупом. Считает сервер",
        ),
        ("ROI по ТЗ", ["—", *[s.roi_tz for s in scenarios[1:]]], SHARE, "Накопленный эффект / вложения"),
        ("Чистый ROI", ["—", *[s.roi_net for s in scenarios[1:]]], SHARE, "(Накопленный эффект - вложения) / вложения"),
        ("NPV", ["—", *[s.npv_rub for s in scenarios[1:]]], MLN, "Чистая приведенная стоимость, считает сервер"),
        (
            "IRR проекта",
            ["—", *[_irr(*project_irr(s)) for s in scenarios[1:]]],
            SHARE,
            "Ставка, при которой NPV = 0, без схемы оплаты. Выше 100% пишем словами: вложения возвращаются быстрее, чем за год",
        ),
    ]
    if any(financed(s) for s in scenarios):
        lines += [
            (
                "IRR своих денег",
                [
                    "—",
                    *[
                        _irr(s.irr, s.npv_rub, s.investment_year0_rub) if financed(s) else "без долга"
                        for s in scenarios[1:]
                    ],
                ],
                SHARE,
                "От первого взноса, с платежами по долгу: при маленьком взносе всегда высокий",
            ),
            (
                "Свои деньги вернутся, лет",
                ["—", *[s.payback_own_years if financed(s) else "без долга" for s in scenarios[1:]]],
                ONE,
                "Только первый взнос; окупаемость по потоку выше считается с учетом остатка долга",
            ),
            (
                "Покрытие долга, минимум по годам",
                ["—", *[_cover(s) if financed(s) else "без долга" for s in scenarios[1:]]],
                ONE,
                "Поток до платежей по долгу / платежи (DSCR). Банк ждет не меньше 1,2 (Альт-Инвест)",
            ),
        ]
    for index, (label, values, fmt, how) in enumerate(lines):
        row = first + index
        shown = [value if value is not None else "не окупается" for value in values]
        _put(sheet, row, [label, *shown, how], fmt, STRONG if index == 0 else TEXT)
        for column in range(2, 2 + count):
            cell = sheet.cell(row=row, column=column)
            cell.alignment = Alignment(horizontal="right", vertical="top")
            if cell.value == "—":
                cell.font = MUTED
        how_cell = sheet.cell(row=row, column=2 + count)
        how_cell.font = MUTED
        how_cell.alignment = Alignment(wrap_text=True, vertical="top")
    last = get_column_letter(2 + count)
    row += 2
    sheet.cell(row=row, column=1, value="Вывод").font = STRONG
    for s in scenarios:
        if s.verdict:
            row += 1
            _paragraph(sheet, row, last, f"{scenario_name(s)}, {s.verdict.band}: {s.verdict.text}")
    row += 2
    sheet.cell(row=row, column=1, value="Что может сдвинуть вывод").font = STRONG
    for index, risk in enumerate(report.risks, 1):
        row += 1
        _paragraph(sheet, row, last, f"{index}. {risk}")


def _paragraph(sheet: Worksheet, row: int, last: str, text: str) -> None:
    """Абзац на всю ширину таблицы. Высоту строки задаем сами: у объединенных ячеек Excel ее не подбирает"""
    sheet.merge_cells(f"A{row}:{last}{row}")
    cell = sheet.cell(row=row, column=1, value=text)
    cell.font = TEXT
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[row].height = 14 * (len(text) // 120 + 1)


def _costs(sheet: Worksheet, report: Report) -> None:
    row = _sheet(sheet, "Затраты по статьям", report, [32, 12, 18, 18])
    purchase, raas = report.scenario("purchase"), report.scenario("raas")
    _head(sheet, row, ["Статья", "Когда", f"{NAMES['purchase']}, ₽", f"{NAMES['raas']}, ₽"])
    for item in purchase.capex_rub:
        row += 1
        values = [CAPEX_NAMES.get(item, item), "разово", purchase.capex_rub.get(item, 0), raas.capex_rub.get(item, 0)]
        _put(sheet, row, values, RUB)
        sheet.cell(row=row, column=2).font = MUTED
    for index in range(len(purchase.years)):
        for item in {**purchase.years[index].opex_rub, **raas.years[index].opex_rub}:
            row += 1
            values = [
                OPEX_NAMES.get(item, item),
                f"{index + 1} год",
                purchase.years[index].opex_rub.get(item, 0),
                raas.years[index].opex_rub.get(item, 0),
            ]
            _put(sheet, row, values, RUB)
            sheet.cell(row=row, column=2).font = MUTED


def _sensitivity(sheet: Worksheet, report: Report) -> None:
    """Что будет, если: на каждый параметр и сдвиг окупаемость и стоимость владения обоих сценариев.
    Числа с сервера: каждый вариант это целый расчет заново, формулой в Excel его не повторить"""
    data = report.sensitivity
    row = _sheet(sheet, "Что будет, если", report, [26, 12, 10, 18, 18, 18, 18, 60], landscape=True)
    sheet["A4"] = "Одно число сдвигаем, остальные как в расчете. При сдвиге объема парк растет вместе со спросом"
    sheet["A4"].font = MUTED
    horizon = data.horizon_years
    _head(
        sheet,
        row,
        [
            "Параметр",
            "Сдвиг",
            "Роботов",
            "Покупка: окупаемость, лет",
            f"Покупка: владение за {horizon} лет, ₽",
            "Аренда: окупаемость, лет",
            f"Аренда: владение за {horizon} лет, ₽",
            "Что меняется",
        ],
    )
    for param in data.params:
        for cell in param.cells:
            row += 1
            found = {o.scenario_id: o for o in cell.outcomes}
            purchase, raas = found.get("purchase"), found.get("raas")
            values = [
                param.name,
                delta(cell.delta),
                cell.fleet,
                _years_or_dash(purchase),
                purchase.tco_rub if purchase else "—",
                _years_or_dash(raas),
                raas.tco_rub if raas else "—",
                f"{param.base}. {param.what}" if cell.delta == data.deltas[0] else "",
            ]
            _put(sheet, row, values, font=STRONG if cell.delta == 0 else TEXT)
            for column in (4, 6):
                sheet.cell(row=row, column=column).number_format = "0.00"
            for column in (5, 7):
                sheet.cell(row=row, column=column).number_format = RUB
            sheet.cell(row=row, column=8).font = MUTED
    if report.sensitivity_all:
        _sensitivity_all(sheet, row + 2, report.sensitivity_all)


def _sensitivity_all(sheet: Worksheet, row: int, data: SensitivityAll) -> None:
    """Все числа по силе влияния, как на экране по кнопке: размах выгоды за горизонт между -20% и +20%"""
    sheet.cell(row=row, column=1, value="Все числа по силе влияния").font = STRONG
    sheet.cell(
        row=row + 1,
        column=1,
        value="Каждое число по одному на -20% и +20%. Размах: насколько меняется выгода против работы без роботов",
    ).font = MUTED
    row += 2
    horizon = data.horizon_years
    _head(
        sheet,
        row,
        [
            "Число",
            "Сейчас",
            "",
            f"Покупка: размах выгоды за {horizon} лет, ₽",
            "Покупка: окупаемость при -20% и +20%",
            f"Аренда: размах выгоды за {horizon} лет, ₽",
            "Аренда: окупаемость при -20% и +20%",
            "Путь в модели",
        ],
    )
    for param in data.params:
        row += 1
        swings = {s.scenario_id: s.rub for s in param.swings}
        _put(
            sheet,
            row,
            [
                param.name,
                param.base,
                "",
                swings.get("purchase", "—"),
                _ends(param, "purchase"),
                swings.get("raas", "—"),
                _ends(param, "raas"),
                param.id,
            ],
        )
        for column in (4, 6):
            sheet.cell(row=row, column=column).number_format = RUB
        sheet.cell(row=row, column=8).font = MUTED
    for title, names_list in (
        ("Итог не двигают при сдвиге до 20%", data.flat),
        ("Равны нулю, в процентах их не сдвинуть", data.zeros),
    ):
        if names_list:
            row += 2
            sheet.cell(row=row, column=1, value=f"{title}: {', '.join(names_list)}").font = MUTED


def _ends(param, scenario_id: str) -> str:
    """Окупаемость при -20% и при +20% одной строкой: 1,8 -> 1,0"""

    def one(cell) -> str:
        if cell.baseline_tco_rub is None:
            return "не справится"
        found = next((o for o in cell.outcomes if o.scenario_id == scenario_id), None)
        if found is None:
            return "—"
        return "не окупается" if found.payback_years is None else f"{found.payback_years:.1f}".replace(".", ",")

    return f"{one(param.cells[0])} -> {one(param.cells[-1])}"


def _years_or_dash(outcome) -> float | str:
    if outcome is None:
        return "—"
    return outcome.payback_years if outcome.payback_years is not None else "не окупается"


def _shift(sheet: Worksheet, report: Report) -> None:
    """Итоги прогона смены, колонка на задачу: каждая задача гоняется своими роботами по тому же плану"""
    shown = [task for task in report.tasks if task.shift]
    row = _sheet(sheet, "Смена на плане", report, [44] + [22] * len(shown))
    sheet["A4"] = "Прогон смены тем парком, что в расчете: пик спроса, очереди и зарядка"
    sheet["A4"].font = MUTED
    _head(sheet, row, ["Что", *[task.title for task in shown]])
    facts = [
        ("Решение", lambda one: one.robot_name, None),
        ("Роботов в прогоне", lambda one: one.shift.fleet, "0"),
        ("Часов смены", lambda one: one.shift.hours, "0"),
        # единица у каждой задачи своя: паллет, строк, м2. Строкой над числами, колонка на задачу
        ("Единица задачи", lambda one: one.unit, None),
        ("Пик спроса с запасом, в час", lambda one: one.shift.design_demand, ONE),
        ("Сделано за смену", lambda one: one.shift.done, "0"),
        ("В среднем за смену, в час", lambda one: one.shift.ops_per_hour, ONE),
        ("Роботы в работе, доля времени", lambda one: one.shift.busy_share, SHARE),
        ("Ждут в очереди, доля времени", lambda one: one.shift.waiting_share, "0.0%"),
        ("На зарядке, доля времени", lambda one: one.shift.charging_share, "0.0%"),
        ("Средний рейс, с", lambda one: one.shift.avg_cycle_s, "0"),
        ("Узкое место", lambda one: one.shift.bottleneck, None),
        ("План", lambda one: "нарисован или по анкете" if one.shift.plan_known else "типовой", None),
    ]
    for label, value, fmt in facts:
        row += 1
        _put(sheet, row, [label, *[value(task) for task in shown]], fmt)


def _tasks(sheet: Worksheet, report: Report) -> None:
    """Разбивка по задачам: строка на задачу, общее на объект и объект целиком, по каждому сценарию.
    Сумма строк равна объекту, ее посчитал сервер"""
    horizon = report.result.horizon_years
    row = _sheet(sheet, "По задачам", report, [40, 30, 10, 18, 18, 18, 18, 16])
    sheet["A4"] = (
        "Общее на объект делается один раз: интеграция с учетной системой, инфраструктура, связь и дежурные "
        "операторы на весь парк"
    )
    sheet["A4"].font = MUTED
    shared = report.result.shared
    for scenario in report.scenarios:
        if scenario.id == "baseline":
            continue
        _head(
            sheet,
            row,
            [
                NAMES[scenario.id],
                "Решение",
                "Роботов",
                "Вложения на старте, ₽",
                "Затраты на роботов, 1 год, ₽",
                "Люди, 1 год, ₽",
                f"Стоимость владения за {horizon} лет, ₽",
                "Окупаемость, лет",
            ],
        )
        for task in report.tasks:
            part = task.scenario(scenario.id)
            if part is None:
                continue
            row += 1
            values = [part.investment_year0_rub, part.opex_year1_rub, part.staff_year1_rub, part.tco_rub]
            payback = part.payback_years if part.payback_years is not None else "не окупается"
            _put(sheet, row, [task.title, task.robot_name, task.part.sizing.fleet, *values, payback], RUB)
            sheet.cell(row=row, column=8).number_format = ONE
            sheet.cell(row=row, column=3).number_format = "0"
        common = next((one for one in shared.scenarios if one.id == scenario.id), None) if shared else None
        if shared and common:
            row += 1
            values = [common.investment_year0_rub, common.opex_year1_rub, common.staff_year1_rub, common.tco_rub]
            _put(sheet, row, ["Общее на объект", f"дежурных в смену {shared.operator_posts}", "", *values, ""], RUB)
        first = scenario.years[0]
        row += 1
        values = [scenario.investment_year0_rub, first.opex_total_rub, first.staff_cost_rub, scenario.tco_rub]
        whole_payback = (
            scenario.payback_cumulative_years if scenario.payback_cumulative_years is not None else "не окупается"
        )
        _put(
            sheet,
            row,
            ["Объект целиком", "", report.result.sizing.fleet, *values, whole_payback],
            RUB,
            font=STRONG,
            border=LAST_ROW,
        )
        sheet.cell(row=row, column=3).number_format = "0"
        sheet.cell(row=row, column=8).number_format = ONE
        row += 2


def _fleet(sheet: Worksheet, report: Report) -> None:
    """Парк: колонка на задачу, у нескольких задач последней колонкой объект целиком"""
    row = _sheet(sheet, "Парк роботов и штат", report, [40, 22, 22, 22, 22])
    sizing = report.result.sizing
    # у одной задачи колонка это объект: у задачи без общего на объект дежурных нет
    columns = [(task.title, task.part.sizing if report.several else sizing, task) for task in report.tasks]
    if report.several:
        columns.append(("Объект целиком", sizing, None))
    _head(sheet, row, ["Что", *[name for name, _, _ in columns]])
    facts = [
        ("Решение", lambda s, task: task.robot_name if task else "", None),
        ("Производитель", lambda s, task: (task.robot.vendor if task.robot else "—") if task else "", None),
        (
            "Цена решения из каталога, ₽",
            lambda s, task: (task.robot.price_rub if task.robot else None) if task else "",
            RUB,
        ),
        ("Откуда размер парка", lambda s, task: s.fleet_source, None),
        ("Роботов", lambda s, task: s.fleet, "0"),
        ("Зарядных станций", lambda s, task: s.chargers, "0"),
        ("Станций комплектации", lambda s, task: s.stations, "0"),
        ("Спрос в пиковый час, операций", lambda s, task: s.design_demand_ops_per_hour, ONE),
        ("Робот делает операций в час", lambda s, task: s.robot_ops_per_hour, ONE),
        ("Средний маршрут, м", lambda s, task: s.route_m, ONE),
        (
            "Дежурных операторов в смену",
            lambda s, task: s.operator_posts if task is None or not report.several else "",
            "0",
        ),
        ("Людей на операции в смену, было", lambda s, task: s.people_per_shift_before, ONE),
        ("Людей на операции в смену, стало", lambda s, task: s.people_per_shift_after, ONE),
        ("Высвобождено ставок", lambda s, task: s.released_fte, ONE),
        (
            "Из них переучиваем в операторов",
            lambda s, task: s.retrained_fte if task is None or not report.several else "",
            ONE,
        ),
    ]
    for label, value, fmt in facts:
        row += 1
        cells = [value(one, task) for _, one, task in columns]
        # спрос, маршрут и производительность у задач в своих единицах, у объекта их нет
        _put(sheet, row, [label, *["—" if cell is None else cell for cell in cells]], fmt)
    row += 2
    for column, name in enumerate(["Роль", "Мест в штате", "Из них работает", "Оклад в месяц, ₽", "По договору"], 1):
        cell = sheet.cell(row=row, column=column, value=name)
        cell.font = HEAD
        cell.border = UNDER_HEAD
        cell.alignment = Alignment(horizontal="left" if column == 1 else "right", wrap_text=True)
    for line in report.staff:
        row += 1
        _put(
            sheet, row, [line.role, line.headcount, line.filled, line.salary_month, "да" if line.contractor else "нет"]
        )
        sheet.cell(row=row, column=4).number_format = RUB


def _parameters(sheet: Worksheet, report: Report) -> None:
    row = _sheet(sheet, "Параметры объекта", report, [40, 14, 14, 10, 12, 64], landscape=True)
    _head(sheet, row, ["Параметр", "Значение", "Единица", "Доверие", "Поправлено", "Источник"])
    for one in report.parameters:
        row += 1
        edited = "вами" if one.path in report.request.overrides else ""
        _put(sheet, row, [one.label, one.value, one.unit, one.trust, edited, one.source])
        sheet.cell(row=row, column=4).font = TRUST
        sheet.cell(row=row, column=5).font = MARK
        sheet.cell(row=row, column=6).font = MUTED
    if report.measures:
        m = report.measures
        row += 2
        sheet.cell(row=row, column=1, value="План объекта").font = STRONG
        for label, value, unit in [
            ("Здание", f"{m.width_m:g} x {m.length_m:g}", "м"),
            ("Средний маршрут до буфера у ворот", round(m.route_m, 1), "м"),
            ("Мест у ворот", m.docks, "шт"),
            ("Проездов между рядами", m.aisles, "шт"),
            ("Самый узкий проезд", m.aisle_m, "м"),
            ("Верхний ярус", m.rack_top_m, "м"),
            ("Пандусов", m.ramps, "шт"),
        ]:
            row += 1
            _put(sheet, row, [label, value, unit])


def _sources(sheet: Worksheet, report: Report) -> None:
    row = _sheet(sheet, "Откуда цифры", report, [22, 40, 14, 18, 9, 70, 10, 44], landscape=True)
    sheet["A4"] = (
        "Доверие: S закон или норматив, A проверено, B сходятся два источника, C один источник, "
        "D только организатор или СМИ, E источники расходятся, F наше допущение"
    )
    sheet["A4"].font = MUTED
    _head(sheet, row, ["Группа", "Значение", "Число", "Единица", "Доверие", "Источник", "Дата", "Путь в модели"])
    order = list(names.GROUPS.values())
    sources = sorted(
        report.result.sources,
        key=lambda one: order.index(names.group(one.path)) if names.group(one.path) in order else 99,
    )
    for source in sources:
        row += 1
        value = source.value if not isinstance(source.value, list) else " - ".join(f"{v:g}" for v in source.value)
        values = [
            names.group(source.path),
            report.source_name(source.path),
            value,
            source.unit,
            source.trust,
            source.source,
            source.date,
            source.path,
        ]
        _put(sheet, row, values)
        sheet.cell(row=row, column=5).font = TRUST
        for column in (1, 4, 6, 7, 8):
            sheet.cell(row=row, column=column).font = MUTED


def _irr(irr: float | None, npv: float | None, investment: float):
    """Ставка числом до 100%, выше словами, как на экране"""
    return irr if irr is not None and irr <= 1 and investment > 0 else irr_text(irr, npv, investment)


def _cover(scenario):
    ratios = [ratio for ratio in scenario.financing.debt_cover if ratio is not None]
    return min(ratios) if ratios else "платежей нет"
